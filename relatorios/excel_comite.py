"""Planilha do comitê mensal: Real x Previsto por ação, KPIs por safra/cluster,
migração de estados e benchmarks.

Fora do núcleo porque depende de openpyxl (pip install openpyxl). As premissas
ficam em células de entrada (azul, fundo amarelo) e tudo o que é cálculo é
fórmula: mudou a premissa, a planilha recalcula ROI e decisão.
"""
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONTE = "Arial"
AZUL = Font(name=FONTE, color="0000FF")
PRETO = Font(name=FONTE)
NEGRITO = Font(name=FONTE, bold=True)
TITULO = Font(name=FONTE, bold=True, size=14)
CAB = Font(name=FONTE, bold=True, color="FFFFFF")
FUNDO_CAB = PatternFill("solid", fgColor="1F3A5F")
FUNDO_PREV = PatternFill("solid", fgColor="DCE6F1")
FUNDO_REAL = PatternFill("solid", fgColor="E2EFDA")
AMARELO = PatternFill("solid", fgColor="FFFF00")
FINO = Border(bottom=Side(style="thin", color="BFBFBF"))
REAIS = 'R$ #,##0.00;(R$ #,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
ROI = '0.00"x";(0.00"x");"-"'
INTEIRO = '#,##0;(#,##0);"-"'
NOMES = {"whatsapp": "WhatsApp", "rcs": "RCS", "agente_voz": "Agente virtual", "discador": "Discador",
         "sms": "SMS", "email": "E-mail", "sem_origem": "Sem origem"}
REGUAS = {"localizacao": "Localização", "giro": "Giro (Não CPC)", "cpc": "CPC / rotação",
          "preventivo": "Preventivo / colchão", "quebra": "Quebra", "outros": "Outros"}


def _cab(ws, linha, textos, fundo=FUNDO_CAB, fonte=CAB):
    for i, t in enumerate(textos, 1):
        c = ws.cell(linha, i, t)
        c.font, c.fill = fonte, fundo
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def _larguras(ws, larguras):
    for i, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def gerar_excel(caminho, periodo: tuple[date, date], realizado: list[dict], kpis: list[dict],
                matriz: dict, bench: list[dict], ordem: list[dict], orcamento: dict, custos_unit: dict,
                ticket_medio: float):
    wb = Workbook()
    _leia_me(wb.active, periodo)
    ws_p = wb.create_sheet("Premissas")
    ref = _premissas(ws_p, orcamento, custos_unit, ticket_medio)
    ws_r = wb.create_sheet("Real x Previsto")
    n = _real_x_previsto(ws_r, realizado, ref)
    _resumo(wb.create_sheet("Resumo", 1), n)
    _kpis(wb.create_sheet("KPIs safra-cluster"), kpis)
    _migracao(wb.create_sheet("Migração de estados"), matriz)
    _benchmarks(wb.create_sheet("Benchmarks"), bench, ordem)
    wb.calculation.fullCalcOnLoad = True  # o Excel recalcula tudo ao abrir
    wb.save(caminho)


def _leia_me(ws, periodo):
    ws.title = "Leia-me"
    ws["A1"], ws["A1"].font = "Comitê de orçamento — Real x Previsto", TITULO
    linhas = [
        f"Período: {periodo[0]:%d/%m/%Y} a {periodo[1]:%d/%m/%Y}. Fonte do realizado: MotorCob (trilha, retornos dos "
        "fornecedores e parcelas do sistema de acordos).",
        "",
        "Como ler",
        "• Texto azul com fundo amarelo = premissa editável (aba Premissas e coluna 'Volume previsto'). Texto preto = fórmula.",
        "• Cada linha de 'Real x Previsto' é uma ação: régua do playbook x canal. O realizado de cada tentativa é "
        "atribuído à régua em que o cliente estava no dia (pela TAG).",
        "• Acordos são atribuídos à régua e ao canal do contato que os originou.",
        "• Recuperação prevista = acordos previstos x ticket médio x % de parcelas pagas. Receita = recuperação x "
        "% de remuneração (comissão/honorários).",
        "• Receita real projetada = valor acordado no período x % de parcelas pagas x % de remuneração (as parcelas "
        "futuras ainda não venceram).",
        "• ROI realizado (caixa) = recuperado até o fim do período x % de remuneração ÷ custo — cresce conforme as "
        "parcelas vencem.",
        "• Custo = custo variável das tentativas nos fornecedores. Custos fixos (operadores, licenças) não estão "
        "incluídos.",
        "• Decisão (critério do playbook): ROI ≥ 1,5x aprovação direta · 1,0x–1,5x aprova com plano · < 1,0x reprova "
        "ou exige justificativa.",
        "• 'Volume previsto' vem pré-preenchido com o volume realizado: assim o previsto mostra o que o benchmark "
        "esperava daquele volume. Para orçar o mês seguinte, substitua pelo volume planejado.",
        "",
        "Abas: Resumo · Premissas · Real x Previsto · KPIs safra-cluster · Migração de estados · Benchmarks "
        "(inclui sugestão de ordem de rotação — só muda por decisão do comitê).",
    ]
    for i, t in enumerate(linhas, 3):
        ws.cell(i, 1, t).font = NEGRITO if t == "Como ler" else PRETO
    ws.column_dimensions["A"].width = 130


def _premissas(ws, orc, custos_unit, ticket):
    ws["A1"], ws["A1"].font = "Premissas", TITULO
    _cab(ws, 3, ["Canal", "Código", "Taxa de contato prevista", "Custo unitário (R$)"])
    ref = {"canal": {}}
    for i, (canal, taxa) in enumerate(orc["taxa_contato_prevista"].items(), 4):
        ws.cell(i, 1, NOMES[canal]).font = PRETO
        ws.cell(i, 2, canal).font = PRETO
        for col, v, fmt in ((3, taxa, PCT), (4, custos_unit.get(canal, 0), REAIS)):
            c = ws.cell(i, col, v)
            c.font, c.fill, c.number_format = AZUL, AMARELO, fmt
        ref["canal"][canal] = i
    ult = 3 + len(orc["taxa_contato_prevista"])
    ws.cell(4, 3).comment = Comment("WhatsApp 15%, RCS 7%, Discador 3%, SMS 1,5%: playbook. Agente virtual e "
                                    "e-mail: premissas a calibrar com o realizado (aba Benchmarks).", "MotorCob")
    ws.cell(4, 4).comment = Comment("Custo médio por tentativa observado no período (retornos dos fornecedores).",
                                    "MotorCob")
    ref["faixa_canal"] = f"Premissas!$B$4:$B${ult}"
    ref["faixa_taxa"] = f"Premissas!$C$4:$C${ult}"
    ref["faixa_custo"] = f"Premissas!$D$4:$D${ult}"
    base = ult + 2
    gerais = [("% acordo sobre contato", orc["pct_acordo_sobre_contato"], PCT, "Premissa: calibrar com o realizado."),
              ("Ticket médio do acordo (R$)", round(ticket, 2), REAIS, "Média do valor acordado no período."),
              ("% de parcelas pagas", orc["pct_parcelas_pagas"], PCT, "Premissa: calibrar com a aba KPIs."),
              ("% de remuneração sobre o recuperado", orc["pct_remuneracao"], PCT,
               "Comissão/honorários da assessoria. Premissa: confirmar no contrato de cada credor."),
              ("ROI para aprovação direta", orc["roi_aprova"], ROI, "Critério do playbook."),
              ("ROI mínimo (aprova com plano)", orc["roi_com_plano"], ROI, "Critério do playbook.")]
    chaves = ["acordo", "ticket", "pagas", "remuneracao", "roi_aprova", "roi_plano"]
    for k, (rotulo, v, fmt, nota) in enumerate(gerais):
        r = base + k
        ws.cell(r, 1, rotulo).font = NEGRITO
        c = ws.cell(r, 3, v)
        c.font, c.fill, c.number_format = AZUL, AMARELO, fmt
        c.comment = Comment(nota, "MotorCob")
        ref[chaves[k]] = f"Premissas!$C${r}"
    _larguras(ws, [32, 14, 22, 20])
    return ref


COLS = ["Régua", "Canal", "Vertical", "Fornecedor",
        # previsto (E:S)
        "Volume previsto", "Custo unitário", "Orçado", "Taxa de contato prevista", "Contatos previstos",
        "% acordo s/ contato", "Acordos previstos", "Ticket médio", "% parcelas pagas", "Recuperação prevista",
        "% remuneração", "Receita prevista", "ROI previsto", "Decisão", "Custo por acordo previsto",
        # real (T:AG)
        "Volume real", "Custo real", "Contatos reais", "Acordos reais", "Valor acordado", "Recuperado (caixa)",
        "Taxa de contato real", "Recuperação real projetada", "Receita real projetada", "ROI real",
        "ROI realizado (caixa)", "Desvio de receita", "Desvio de custo", "Custo por acordo real"]
N_COLS = len(COLS)          # 33; a coluna 34 (AH, oculta) guarda o código do canal
INI_REAL = COLS.index("Volume real") + 1


def _decisao(celula_roi, ref):
    return (f'=IF({celula_roi}>={ref["roi_aprova"]},"Aprovação direta",IF({celula_roi}>={ref["roi_plano"]},'
            f'"Aprova com plano","Reprova ou justificar"))')


def _real_x_previsto(ws, realizado, ref):
    ws["A1"], ws["A1"].font = "Real x Previsto por ação (régua x canal)", TITULO
    ws.cell(2, 5, "PREVISTO").font = NEGRITO
    ws.cell(2, INI_REAL, "REAL (MotorCob)").font = NEGRITO
    _cab(ws, 3, COLS)
    for col in range(5, INI_REAL):
        ws.cell(3, col).fill = PatternFill("solid", fgColor="2E5A88")
    for col in range(INI_REAL, N_COLS + 1):
        ws.cell(3, col).fill = PatternFill("solid", fgColor="375623")
    ini = 4
    fmts = {6: REAIS, 7: REAIS, 8: PCT, 9: '#,##0.0', 10: PCT, 11: '#,##0.0', 12: REAIS, 13: PCT, 14: REAIS,
            15: PCT, 16: REAIS, 17: ROI, 19: REAIS, 20: INTEIRO, 21: REAIS, 22: INTEIRO, 23: INTEIRO, 24: REAIS,
            25: REAIS, 26: PCT, 27: REAIS, 28: REAIS, 29: ROI, 30: ROI, 31: PCT, 32: PCT, 33: REAIS}
    for i, r in enumerate(realizado):
        L = ini + i
        canal = r["canal"]
        for col, v in enumerate([REGUAS.get(r["regua"], r["regua"]), NOMES.get(canal, canal), r["vertical"],
                                 r["fornecedor"]], 1):
            ws.cell(L, col, v).font = PRETO
        ws.cell(L, N_COLS + 1, canal)  # chave técnica das fórmulas (coluna oculta)
        c = ws.cell(L, 5, r["volume"])
        c.font, c.fill = AZUL, AMARELO
        busca = lambda faixa: f'=IFERROR(INDEX({ref[faixa]},MATCH($AH{L},{ref["faixa_canal"]},0)),0)'  # noqa: E731
        f = {6: busca("faixa_custo"), 7: f"=E{L}*F{L}", 8: busca("faixa_taxa"), 9: f"=E{L}*H{L}",
             10: f"={ref['acordo']}", 11: f"=I{L}*J{L}", 12: f"={ref['ticket']}", 13: f"={ref['pagas']}",
             14: f"=K{L}*L{L}*M{L}", 15: f"={ref['remuneracao']}", 16: f"=N{L}*O{L}",
             17: f"=IF(G{L}>0,P{L}/G{L},0)", 18: _decisao(f"Q{L}", ref), 19: f"=IF(K{L}>0,G{L}/K{L},0)",
             26: f"=IF(T{L}>0,V{L}/T{L},0)", 27: f"=X{L}*M{L}", 28: f"=AA{L}*O{L}",
             29: f"=IF(U{L}>0,AB{L}/U{L},0)", 30: f"=IF(U{L}>0,Y{L}*O{L}/U{L},0)",
             31: f"=IF(P{L}>0,AB{L}/P{L}-1,0)", 32: f"=IF(G{L}>0,U{L}/G{L}-1,0)", 33: f"=IF(W{L}>0,U{L}/W{L},0)"}
        for col, v in f.items():
            ws.cell(L, col, v).font = PRETO
        for col, k in ((20, "volume"), (21, "custo"), (22, "contatos"), (23, "acordos"),
                       (24, "valor_acordado"), (25, "recebido")):
            ws.cell(L, col, r[k]).font = PRETO
        for col in range(1, N_COLS + 1):
            ws.cell(L, col).border = FINO
            if col in fmts:
                ws.cell(L, col).number_format = fmts[col]
        ws.cell(L, 5).number_format = INTEIRO
    fim = ini + len(realizado) - 1
    T = fim + 1
    ws.cell(T, 1, "TOTAL")
    for col in (5, 7, 9, 11, 14, 16, 20, 21, 22, 23, 24, 25, 27, 28):
        letra = get_column_letter(col)
        ws.cell(T, col, f"=SUM({letra}{ini}:{letra}{fim})")
    tot = {17: f"=IF(G{T}>0,P{T}/G{T},0)", 18: _decisao(f"Q{T}", ref), 19: f"=IF(K{T}>0,G{T}/K{T},0)",
           26: f"=IF(T{T}>0,V{T}/T{T},0)", 29: f"=IF(U{T}>0,AB{T}/U{T},0)",
           30: f"=IF(U{T}>0,Y{T}*{ref['remuneracao']}/U{T},0)", 31: f"=IF(P{T}>0,AB{T}/P{T}-1,0)",
           32: f"=IF(G{T}>0,U{T}/G{T}-1,0)", 33: f"=IF(W{T}>0,U{T}/W{T},0)"}
    for col, v in tot.items():
        ws.cell(T, col, v)
    for col in range(1, N_COLS + 1):
        c = ws.cell(T, col)
        c.font = NEGRITO
        c.number_format = fmts.get(col, INTEIRO if col == 5 else "General")
    ws.cell(3, 5).comment = Comment("Pré-preenchido com o volume realizado. Para orçar o próximo mês, "
                                    "substitua pelo volume planejado.", "MotorCob")
    ws.cell(3, 21).comment = Comment("Custo variável das tentativas (fornecedores). Custos fixos (operadores, "
                                     "licenças) não estão incluídos.", "MotorCob")
    ws.column_dimensions[get_column_letter(N_COLS + 1)].hidden = True
    ws.freeze_panes = "E4"
    ws.row_dimensions[3].height = 45
    _larguras(ws, [20, 14, 9, 12] + [13] * (N_COLS - 4))
    return {"ini": ini, "fim": fim}


def _resumo(ws, n):
    ws["A1"], ws["A1"].font = "Resumo do período", TITULO
    _cab(ws, 3, ["", "Orçado (previsto)", "Custo real", "Receita prevista", "Receita real projetada",
                 "Recuperado (caixa)", "ROI previsto", "ROI real", "Acordos previstos", "Acordos reais"])
    faixa = lambda col: f"'Real x Previsto'!${col}${n['ini']}:${col}${n['fim']}"  # noqa: E731
    for i, vert in enumerate(["Voz", "Digital", "Total"], 4):
        ws.cell(i, 1, vert).font = NEGRITO
        for col, letra, fmt in ((2, "G", REAIS), (3, "U", REAIS), (4, "P", REAIS), (5, "AB", REAIS),
                                (6, "Y", REAIS), (9, "K", '#,##0.0'), (10, "W", INTEIRO)):
            formula = f"=SUM({faixa(letra)})" if vert == "Total" else \
                f'=SUMIFS({faixa(letra)},{faixa("C")},"{vert}")'
            c = ws.cell(i, col, formula)
            c.number_format, c.font = fmt, PRETO
        ws.cell(i, 7, f"=IF(B{i}>0,D{i}/B{i},0)").number_format = ROI
        ws.cell(i, 8, f"=IF(C{i}>0,E{i}/C{i},0)").number_format = ROI
    ws.cell(8, 1, "Decisão por ação (ROI previsto)").font = NEGRITO
    for i, d in enumerate(["Aprovação direta", "Aprova com plano", "Reprova ou justificar"], 9):
        ws.cell(i, 1, d).font = PRETO
        ws.cell(i, 2, f'=COUNTIF({faixa("R")},"{d}")').number_format = INTEIRO
    ws.cell(13, 1, "Custo = custo variável das tentativas nos fornecedores; custos fixos não estão incluídos, "
                   "por isso o ROI por ação tende a ser alto. Receita = recuperação × % de remuneração.").font = PRETO
    _larguras(ws, [30] + [18] * 9)


def _kpis(ws, kpis):
    ws["A1"], ws["A1"].font = "KPIs por safra e cluster (frentes do playbook)", TITULO
    if not kpis:
        return
    cols = list(kpis[0].keys())
    _cab(ws, 3, [c.replace("pct_", "% ").replace("_", " ") for c in cols])
    for i, k in enumerate(kpis, 4):
        for j, c in enumerate(cols, 1):
            cel = ws.cell(i, j, k[c])
            cel.font = NEGRITO if k["safra"] == "TOTAL" else PRETO
            if c.startswith("pct_"):
                cel.number_format = PCT
            elif c.startswith("custo"):
                cel.number_format = REAIS
    ws.freeze_panes = "C4"
    ws.row_dimensions[3].height = 45
    _larguras(ws, [10, 8] + [12] * (len(cols) - 2))


def _migracao(ws, matriz):
    ws["A1"], ws["A1"].font = "Migração de estados no período (de → para)", TITULO
    estados = ["ENTRADA", "LOC", "CPA", "CPB", "NCP", "PRE", "COL", "QBR", "LIQ", "BLQ"]
    _cab(ws, 3, ["de \\ para"] + estados[1:] + ["Total"])
    for i, a in enumerate(estados, 4):
        ws.cell(i, 1, a).font = NEGRITO
        for j, b in enumerate(estados[1:], 2):
            c = ws.cell(i, j, matriz.get(a, {}).get(b, 0))
            c.number_format, c.font = INTEIRO, PRETO
        ult = get_column_letter(len(estados))
        ws.cell(i, len(estados) + 1, f"=SUM(B{i}:{ult}{i})").number_format = INTEIRO
    _larguras(ws, [12] + [9] * len(estados))


def _benchmarks(ws, bench, ordem):
    ws["A1"], ws["A1"].font = "Benchmarks observados (recalibram as premissas do mês seguinte)", TITULO
    cols = ["Régua", "Canal", "Tentativas", "Contatos", "Custo", "Taxa de contato", "Custo por contato",
            "% acordo s/ contato"]
    _cab(ws, 3, cols)
    for i, b in enumerate(bench, 4):
        vals = [REGUAS.get(b["regua"], b["regua"]), NOMES.get(b["canal"], b["canal"]), b["volume"], b["contatos"],
                b["custo"], b["taxa_contato"], b["custo_por_contato"], b["pct_acordo_sobre_contato"]]
        for j, (v, fmt) in enumerate(zip(vals, [None, None, INTEIRO, INTEIRO, REAIS, PCT, REAIS, PCT]), 1):
            c = ws.cell(i, j, v)
            c.font = PRETO
            if fmt:
                c.number_format = fmt
    r0 = 5 + len(bench)
    ws.cell(r0, 1, "Sugestão de ordem de rotação (menor custo por contato nas réguas massivas)").font = NEGRITO
    ws.cell(r0 + 1, 1, "Sugestão para o comitê: a ordem vigente em regras/regua.json só muda por decisão do comitê.")
    _cab(ws, r0 + 2, ["Posição sugerida", "Canal", "Posição atual", "Tentativas", "Contatos", "Taxa de contato",
                      "Custo por contato"])
    for i, o in enumerate(ordem, r0 + 3):
        vals = [o["posicao_sugerida"], NOMES.get(o["canal"], o["canal"]), o["posicao_atual"], o["volume"],
                o["contatos"], o["taxa_contato"], o["custo_por_contato"]]
        for j, (v, fmt) in enumerate(zip(vals, [None, None, None, INTEIRO, INTEIRO, PCT, REAIS]), 1):
            c = ws.cell(i, j, v)
            c.font = NEGRITO if o["posicao_sugerida"] != o["posicao_atual"] else PRETO
            if fmt:
                c.number_format = fmt
    _larguras(ws, [22, 16, 12, 12, 14, 14, 16, 16])
