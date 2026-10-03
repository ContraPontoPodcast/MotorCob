"""Sincroniza a rotina do motor com o Supabase (site motorcob.online).

    python -m nuvem.sincronizar dia    [--data AAAA-MM-DD] [--empresa slug] [--dados ~/MotorCob-dados]
    python -m nuvem.sincronizar comite --mes AAAA-MM   [--empresa slug] [--dados ~/MotorCob-dados]
    python -m nuvem.sincronizar vigiar [--checar]       (roda o dia só de quem tem carga nova no site)

Roda para cada empresa ativa em public.empresas (ou só a de --empresa). Cada empresa tem
a própria pasta de dados, <dados>/empresas/<slug>/, e o próprio arquivo de entrada no
repositório, empresas/<slug>.json (como ler a base bruta e a ocorrência dela). Uma
empresa com erro não impede as outras.

dia, para cada empresa:
  1. baixa os arquivos que o site recebeu (envios pendentes da empresa): base bruta vai
     para bruto/ e vira base/clientes.csv e base/contatos.csv; ocorrências vão para
     ocorrencias/; retornos de fornecedor para retornos/ com o nome original (é pelo
     nome que o motor acha o layout); clientes/contatos/parcelas já no formato do motor
     substituem base/; portal vira logs/portal.csv;
  2. roda a rotina diária (rodar_dia.py);
  3. publica: estado_cliente (upsert), trilha (só o que ainda não subiu), fila_dia do dia
     (só IDs), arquivos de saida/<data>/ em saidas/<slug>/<data>/, envios processados com
     o relatório de ingestão, e a execução com resumo e alertas.
  Se algo falhar, a execução da empresa fica com status 'erro' e os envios dela
  continuam pendentes.

vigiar: para o Mac rodar a cada 2 minutos. Se uma empresa subiu carga do credor no site
  (envio tipo base pendente), roda o dia dela na hora e a lista sai no site em minutos.
  Sem carga nova, não faz nada. Se a rodada der erro, só tenta de novo quando chegar outra
  carga (não fica repetindo o mesmo erro). --checar: só diz se há trabalho (saída 0) ou não (3).

comite: roda o relatório do mês de cada empresa, sobe em saidas/<slug>/comite/<mês>/ e
grava kpis.

Configuração (nunca no git): variáveis MOTORCOB_SUPABASE_URL e MOTORCOB_SUPABASE_KEY
(chave service_role) ou o arquivo <dados>/config/supabase.env criado por
scripts/configurar_nuvem.sh.
"""
import argparse
import calendar
import csv
import json
import os
import re
import sys
import traceback
from collections import Counter
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

from nuvem.supabase_api import Supabase  # noqa: E402

DESTINO_BASE = {"clientes": "base/clientes.csv", "contatos": "base/contatos.csv",
                "parcelas": "base/parcelas.csv", "portal": "logs/portal.csv"}
PASTA_DO_TIPO = {"retorno": "retornos", "base": "bruto", "ocorrencia": "ocorrencias",
                 "enriquecimento": "enriquecimento", "incremental": "incremental", "retirada": "retirada",
                 "acordo": "acordo", "baixa": "baixa"}
TIPOS_CARTEIRA = ("base", "incremental", "retirada", "acordo", "baixa")   # chegou um: a vigia roda o credor
PASTAS_CARTEIRA = ("bruto", "incremental", "retirada", "acordo", "baixa")
MARCADOR_TRILHA = "nuvem_trilha_enviada.txt"
PASTA_EMPRESAS = RAIZ / "empresas"


def agora():
    return datetime.now(timezone.utc).isoformat()


def carregar_config(dados: Path) -> tuple[str, str]:
    url, chave = os.environ.get("MOTORCOB_SUPABASE_URL"), os.environ.get("MOTORCOB_SUPABASE_KEY")
    arq = dados / "config" / "supabase.env"
    if (not url or not chave) and arq.exists():
        for linha in arq.read_text(encoding="utf-8").splitlines():
            k, _, v = linha.partition("=")
            if k.strip() == "MOTORCOB_SUPABASE_URL":
                url = url or v.strip()
            elif k.strip() == "MOTORCOB_SUPABASE_KEY":
                chave = chave or v.strip()
    if not url or not chave:
        raise SystemExit(f"Supabase não configurado: rode scripts/configurar_nuvem.sh (ou crie {arq})")
    return url, chave


def nome_seguro(nome: str) -> str:
    """Só o nome do arquivo, sem pastas (evita gravar fora de retornos/)."""
    nome = Path(nome.replace("\\", "/")).name
    return re.sub(r"[^A-Za-z0-9._-]", "_", nome) or "arquivo.csv"


def empresas_ativas(sb: Supabase, slug: str | None = None) -> list[dict]:
    filtros = {"ativa": "eq.true"}
    if slug:
        filtros["slug"] = f"eq.{slug}"
    lista = sb.selecionar("empresas", filtros, ordem="slug.asc") or []
    if slug and not lista:
        raise SystemExit(f"empresa '{slug}' não existe ou não está ativa no site")
    return lista


def pasta_empresa(dados: Path, emp: dict) -> Path:
    if not re.fullmatch(r"[a-z0-9][a-z0-9-]{1,39}", emp["slug"]):
        raise ValueError(f"slug inválido: {emp['slug']!r}")
    return dados / "empresas" / emp["slug"]


def credores(sb: Supabase, dados: Path, emp: dict) -> list[dict]:
    """Credores ativos da empresa, cada um com a sua pasta: <empresa>/credores/<codigo>/.
    Banco sem a tabela de credores (ou empresa sem credor): uma unidade só, na pasta da empresa
    (id None), como antes. Na primeira vez do credor 'principal', o que estava na pasta da
    empresa (estado, base, saídas…) passa para a pasta dele."""
    base = pasta_empresa(dados, emp)
    try:
        lista = sb.selecionar("credores", {"empresa_id": f"eq.{emp['id']}", "ativo": "eq.true"}, ordem="id.asc") or []
    except Exception:  # noqa: BLE001 — banco ainda sem a migração de credores
        lista = []
    if not lista:
        return [{"id": None, "codigo": None, "nome": emp.get("nome"), "estrategia_id": None, "pasta": base}]
    saida = []
    for c in lista:
        if not re.fullmatch(r"[a-z0-9][a-z0-9-]{0,39}", c["codigo"]):
            raise ValueError(f"código de credor inválido: {c['codigo']!r}")
        pasta = base / "credores" / c["codigo"]
        if c["codigo"] == "principal" and not pasta.exists() and (base / "estado").exists():
            pasta.mkdir(parents=True)
            for nome in ("estado", "base", "saida", "retornos", "ocorrencias", "enriquecimento", "config", "logs",
                         "acoes") + PASTAS_CARTEIRA:
                if (base / nome).exists():
                    (base / nome).replace(pasta / nome)
        saida.append({**c, "pasta": pasta})
    return saida


def arquivo_entrada(emp: dict, pasta_empresas: Path = PASTA_EMPRESAS) -> Path | None:
    arq = pasta_empresas / f"{emp['slug']}.json"
    return arq if arq.exists() else None


# ------------------------------------------------------------------ dia
def planilha_para_csv(conteudo: bytes) -> bytes:
    """Excel (.xlsx) -> CSV com ';' (1ª aba). Datas em dd/mm/aaaa; o resto como está na célula."""
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    saida = io.StringIO()
    w = csv.writer(saida, delimiter=";")
    for linha in ws.iter_rows(values_only=True):
        if linha is None or all(c is None or str(c).strip() == "" for c in linha):
            continue
        w.writerow(["" if c is None else c.strftime("%d/%m/%Y") if hasattr(c, "strftime")
                    else (str(int(c)) if isinstance(c, float) and c.is_integer() else str(c)) for c in linha])
    return saida.getvalue().encode("utf-8")


def baixar_entradas(sb: Supabase, dados: Path, empresa_id, out=print, unidades: list[dict] | None = None,
                    so_credores: set | None = None):
    """Baixa os envios pendentes. Com credores, cada arquivo vai para a pasta do credor dele; o
    retorno do bureau sem credor vai para todos. Devolve [(envio, destino)] ou, com unidades,
    {credor_id: [(envio, destino)]}. so_credores: só os envios desses credores (vigia)."""
    pendentes = sb.selecionar("envios", {"status": "eq.pendente", "empresa_id": f"eq.{empresa_id}"},
                              ordem="enviado_em.asc,id.asc") or []
    por_id = {u["id"]: u for u in unidades or []}
    baixados = {u["id"]: [] for u in unidades or []}
    for e in pendentes:
        cid = e.get("credor_id")
        if unidades is not None:
            if so_credores is not None and cid not in so_credores:
                continue
            if None in por_id:
                alvos = [por_id[None]]
            elif cid is None and e["tipo"] == "enriquecimento":
                alvos = list(por_id.values())
            elif cid in por_id:
                alvos = [por_id[cid]]
            else:
                sb.atualizar("envios", {"id": f"eq.{e['id']}"},
                             {"status": "erro", "relatorio": {"erro": "credor inativo ou não informado"}})
                continue
        else:
            alvos = [{"id": None, "pasta": dados}]
        try:
            conteudo = sb.baixar("entradas", e["caminho"])
        except Exception as ex:  # noqa: BLE001 — um arquivo ruim não derruba os outros
            sb.atualizar("envios", {"id": f"eq.{e['id']}"},
                         {"status": "erro", "relatorio": {"erro": f"não foi possível baixar: {ex}"}})
            out(f"  envio {e['id']} ({e['nome_original']}): erro ao baixar")
            continue
        nome = e["nome_original"]
        if nome.lower().endswith((".xlsx", ".xlsm")):
            try:
                conteudo, nome = planilha_para_csv(conteudo), nome.rsplit(".", 1)[0] + ".csv"
            except Exception as ex:  # noqa: BLE001
                sb.atualizar("envios", {"id": f"eq.{e['id']}"},
                             {"status": "erro", "relatorio": {"erro": f"não consegui abrir a planilha: {ex}"}})
                continue
        elif nome.lower().endswith(".xls"):
            sb.atualizar("envios", {"id": f"eq.{e['id']}"}, {"status": "erro", "relatorio": {
                "erro": "planilha no formato antigo (.xls): salve como .xlsx ou .csv e envie de novo"}})
            continue
        for u in alvos:
            if e["tipo"] in PASTA_DO_TIPO:
                destino = u["pasta"] / PASTA_DO_TIPO[e["tipo"]] / nome_seguro(nome)
            else:
                destino = u["pasta"] / DESTINO_BASE[e["tipo"]]
            destino.parent.mkdir(parents=True, exist_ok=True)
            destino.write_bytes(conteudo)
            if unidades is None:
                baixados.setdefault(None, []).append((e, destino))
            else:
                baixados[u["id"]].append((e, destino))
    n = sum(len(v) for v in baixados.values())
    out(f"  entradas: {n} arquivo(s) baixado(s) do site")
    return baixados if unidades is not None else baixados.get(None, [])


def _c(cid) -> dict:
    """Coluna credor_id nas linhas (sem credor = banco antigo, sem a coluna)."""
    return {"credor_id": cid} if cid is not None else {}


def _fc(cid) -> dict:
    return {"credor_id": f"eq.{cid}"} if cid is not None else {}


def _linhas_estado(estados, empresa_id, clientes=None, dia=None, cid=None):
    ts = agora()
    clientes = clientes or {}

    def extra(k):
        c = clientes.get(k)
        if c is None:
            return {}
        return {"saldo": round(c.saldo, 2), "dias_atraso": c.atraso_em(dia) if dia else c.dias_atraso}

    return [{"empresa_id": empresa_id, **_c(cid), "id_cliente": k, "tag": e.tag, "safra": e.safra.isoformat(),
             "cluster_origem": e.cluster_origem, "cluster_atual": e.cluster_atual, "estado": e.estado,
             "canal": e.canal, "ciclo": e.ciclo, "reenriquecer": e.reenriquecer, "atualizado_em": ts, **extra(k)}
            for k, e in estados.items()]


def _linhas_fila(fila, data, empresa_id, cid=None):
    vistos, saida = set(), []
    for l in fila:
        chave = (l["canal"], l["id_cliente"], bool(l["condicao"]))
        if chave in vistos:
            continue
        vistos.add(chave)
        saida.append({"empresa_id": empresa_id, **_c(cid), "data": data.isoformat(), "canal": l["canal"],
                      "id_cliente": l["id_cliente"], "reserva": bool(l["condicao"]), "regua": l["regua"],
                      "passo": l["passo"], "tag": l["tag"]})
    return saida


def publicar_trilha(sb: Supabase, pasta_estado: Path, empresa_id, cid=None) -> int:
    """Sobe as linhas de estado/trilha.csv que ainda não foram enviadas (retomável)."""
    arq, marca = pasta_estado / "trilha.csv", pasta_estado / MARCADOR_TRILHA
    if not arq.exists():
        return 0
    with open(arq, newline="", encoding="utf-8") as f:
        linhas = list(csv.DictReader(f))
    ja = int(marca.read_text()) if marca.exists() else 0
    novas = linhas[ja:]
    for i in range(0, len(novas), 1000):
        lote = novas[i:i + 1000]
        sb.inserir("trilha", [{"empresa_id": empresa_id, **_c(cid), **{k: l[k] for k in (
            "id_cliente", "data", "tag_anterior", "tag", "motivo", "quem_marcou")}} for l in lote])
        marca.write_text(str(ja + i + len(lote)))  # avança só depois de cada lote confirmado
    return len(novas)


def publicar_arquivos(sb: Supabase, pasta: Path, prefixo: str) -> int:
    n = 0
    for arq in sorted(p for p in pasta.rglob("*") if p.is_file()):
        rel = arq.relative_to(pasta).as_posix()
        tipo = {"xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "txt": "text/plain; charset=utf-8"}.get(arq.suffix.lstrip("."), "text/csv; charset=utf-8")
        sb.enviar("saidas", f"{prefixo}/{rel}", arq.read_bytes(), tipo)
        n += 1
    return n


def _relatorio_envio(e, destino, r, rel_base=None, rel_enriq=None):
    nome = destino.name
    if e["tipo"] == "enriquecimento":
        if rel_enriq is None:
            return "erro", {"erro": "retorno de enriquecimento sem base da empresa ou sem layout em empresas/<slug>.json"}
        if nome in rel_enriq["sem_layout"]:
            return "erro", {"erro": "arquivo não reconhecido como retorno de enriquecimento: confira o nome"}
        arq = next((a for a in rel_enriq["arquivos"] if a["arquivo"] == nome), None)
        return ("processado", arq) if arq else ("processado", {"observacao": "arquivo recebido"})
    if e["tipo"] in TIPOS_CARTEIRA:
        arq = next((a for a in (rel_base or {}).get("arquivos", []) if a["arquivo"] == nome), None)
        if arq is None:
            if e["tipo"] != "base":
                return "erro", {"erro": "arquivo recebido, mas o credor ainda não tem carga geral: envie a carga"}
            return "erro", {"erro": "arquivo não reconhecido como base bruta: confira o nome do arquivo"}
        if e["tipo"] in ("retirada", "acordo", "baixa"):
            return "processado", {**arq, "retirados": rel_base.get("retirados") or {},
                                  "pagamentos": rel_base.get("pagamentos") or {},
                                  "parcelas_pagas": rel_base.get("parcelas_pagas", 0),
                                  "acordos_abertos": rel_base.get("acordos_abertos", 0),
                                  "avisos": rel_base.get("avisos") or {}, "na_carga": rel_base["na_carga"]}
        return "processado", {**arq, "clientes_na_base": rel_base["clientes"], "contatos": rel_base["contatos"],
                              "rejeitadas": rel_base["rejeitadas"], "na_carga": rel_base["na_carga"],
                              "fora_da_carga": rel_base["fora_da_carga"]}
    if e["tipo"] == "ocorrencia":
        if nome in r["sem_layout"] or any(s.startswith(nome + " ") for s in r["sem_layout"]):
            return "erro", {"erro": "arquivo não reconhecido como ocorrência: confira o nome do arquivo"}
        rel = next((x for x in r["relatorios_ocorrencia"] if x.arquivo == nome), None)
        if rel is None:
            return "processado", {"observacao": "arquivo recebido"}
        return "processado", {"linhas": rel.linhas, "aceitas": rel.aceitas, "duplicadas": rel.duplicadas,
                              "contato_identificado": rel.contato_identificado, "avisos": dict(rel.avisos),
                              "rejeitadas": dict(rel.rejeitadas), "quarentena": dict(rel.desconhecidos)}
    if e["tipo"] != "retorno":
        with open(destino, encoding="utf-8", errors="replace") as f:
            return "processado", {"linhas": max(sum(1 for _ in f) - 1, 0)}
    if nome in r["sem_layout"] or any(s.startswith(nome + " ") for s in r["sem_layout"]):
        return "erro", {"erro": "arquivo sem layout de fornecedor: crie layouts/<fornecedor>.json no motor"}
    rel = next((x for x in r["relatorios"] if x.arquivo == nome), None)
    if rel is None:
        return "processado", {"observacao": "arquivo recebido"}
    return "processado", {"fornecedor": rel.fornecedor, "linhas": rel.linhas, "aceitas": rel.aceitas,
                          "duplicadas": rel.duplicadas, "rejeitadas": dict(rel.rejeitadas),
                          "quarentena": dict(rel.desconhecidos)}


def _baixar(sb: Supabase, pasta: Path, tabela: str, empresa_id, ordem: str) -> list[dict]:
    linhas = sb.selecionar(tabela, {"empresa_id": f"eq.{empresa_id}"}, ordem=ordem) or []
    (pasta / "config").mkdir(parents=True, exist_ok=True)
    (pasta / "config" / f"{tabela}.json").write_text(json.dumps(linhas, ensure_ascii=False, indent=1),
                                                     encoding="utf-8")
    return linhas


def chave_sugestao(s: dict) -> str:
    return f'{s["persona"]}|{s["fase"]}|{s["dia"]}|{s["de"]}|{s["para"]}'


def publicar_personas(sb: Supabase, empresa_id, r, cid=None) -> int:
    """Personas do dia (substitui as anteriores) e sugestões novas; devolve quantas sugestões novas."""
    ts = agora()
    sb.apagar("personas", {"empresa_id": f"eq.{empresa_id}", **_fc(cid)})
    linhas = [{"empresa_id": empresa_id, **_c(cid), "persona": p["persona"], "nome": p["nome"], "clientes": p["clientes"],
               "condicoes": p["condicoes"], "ranking": p["ranking"],
               "caracteristicas": r.get("caracteristicas_persona") or [], "atualizado_em": ts}
              for p in r.get("personas") or []]
    if linhas:
        sb.inserir("personas", linhas)
    existentes = {s["chave"]: s for s in sb.selecionar("sugestoes", {"empresa_id": f"eq.{empresa_id}", **_fc(cid)})
                  or []}
    novas = 0
    for s in r.get("sugestoes") or []:
        chave = chave_sugestao(s)
        if chave not in existentes:
            sb.inserir("sugestoes", {"empresa_id": empresa_id, **_c(cid), "chave": chave, "texto": s["texto"], "dados": s})
            novas += 1
        elif existentes[chave]["status"] == "pendente":  # evidência do dia
            sb.atualizar("sugestoes", {"id": f"eq.{existentes[chave]['id']}"}, {"texto": s["texto"], "dados": s})
    return novas


def _acoes_de(linhas, dia) -> dict:
    from motor.acoes import totais
    return totais([l for l in linhas or [] if l["data"] == dia.isoformat()])


def publicar_acoes(sb: Supabase, empresa_id, r, cid=None) -> int:
    """Ações realizadas (só totais): regrava a janela que a rotina recalculou."""
    linhas = r.get("acoes") or []
    desde = min((l["data"] for l in linhas), default=None)
    if desde is None:
        return 0
    sb.apagar("acoes_dia", {"empresa_id": f"eq.{empresa_id}", **_fc(cid), "data": f"gte.{desde}"})
    ts = agora()
    sb.inserir("acoes_dia", [{**l, "empresa_id": empresa_id, **_c(cid), "atualizado_em": ts} for l in linhas])
    return len(linhas)


def aplicar_sugestoes(sb: Supabase, empresa_id, out=print, cid=None) -> int:
    """Sugestões aprovadas no site viram um segmento (condições da persona, no topo da lista) com
    uma estratégia igual à que a persona seguia, com a troca de canal sugerida."""
    import copy
    from motor.persona import condicoes_cluster
    from motor.regua import carregar_regua
    aprovadas = sb.selecionar("sugestoes", {"empresa_id": f"eq.{empresa_id}", **_fc(cid),
                                            "status": "eq.aprovada"}) or []
    if not aprovadas:
        return 0
    clusters = sb.selecionar("clusters", {"empresa_id": f"eq.{empresa_id}"}) or []
    usados = {c["codigo"] for c in clusters}
    ordem = min([c.get("ordem") or 100 for c in clusters] + [100]) - 10
    for sug in aprovadas:
        d = sug["dados"]
        base = {}
        if d.get("estrategia_base") is not None:
            achou = sb.selecionar("estrategias", {"id": f"eq.{d['estrategia_base']}"}) or []
            base = copy.deepcopy(achou[0]["definicao"]) if achou else {}
        loc = base.setdefault("localizacao", {})
        passos = loc.get("passos") or copy.deepcopy(carregar_regua()["localizacao"]["passos"])
        passo = passos.get(str(d["dia"])) or [d["de"]]
        primeiro = passo[0]
        passo[0] = d["para"] if isinstance(primeiro, str) else {**primeiro, "canal": d["para"]}
        passos[str(d["dia"])] = passo
        loc["passos"] = passos
        nome = f"Persona {d['nome']}"[:70] + f" #{sug['id']}"
        estr = sb.inserir("estrategias", {"empresa_id": empresa_id, "nome": nome,
                                          "descricao": sug["texto"], "definicao": base}, retornar=True)[0]
        codigo = next(f"P{n}" for n in range(1, 1000) if f"P{n}" not in usados)
        usados.add(codigo)
        sb.inserir("clusters", {"empresa_id": empresa_id, **_c(cid), "ordem": ordem, "codigo": codigo, "ativo": True,
                                "nome": f"Persona {d['nome']}"[:120],
                                "condicoes": condicoes_cluster(d["condicoes"]), "estrategia_id": estr["id"]})
        ordem -= 10
        sb.atualizar("sugestoes", {"id": f"eq.{sug['id']}"}, {"status": "aplicada", "aplicada_em": agora()})
        out(f"  sugestão aplicada: {sug['texto']} (segmento {codigo})")
    return len(aprovadas)


def baixar_clusters(sb: Supabase, pasta: Path, empresa_id, cid=None) -> list[dict]:
    """Regras de cluster ativas do credor (as dele e as que valem para todos os credores);
    cópia em config/clusters.json (auditoria da rodada)."""
    regras = sb.selecionar("clusters", {"empresa_id": f"eq.{empresa_id}", "ativo": "eq.true"},
                           ordem="ordem.asc,codigo.asc") or []
    regras = [c for c in regras if c.get("credor_id") in (None, cid)]
    (pasta / "config").mkdir(parents=True, exist_ok=True)
    (pasta / "config" / "clusters.json").write_text(json.dumps(regras, ensure_ascii=False, indent=1),
                                                    encoding="utf-8")
    return regras


def _separar_cargas_ruins(sb: Supabase, entrada, pasta: Path, baixados, out=print):
    """Arquivo da carteira com coluna obrigatória faltando (ou sem layout configurado) não entra
    na pasta (senão travaria as próximas rodadas): vai para <tipo>/rejeitados/ e o envio fica com
    erro e o motivo. Se todas as cargas novas (geral/incremental) forem rejeitadas, a rodada para
    (não gera lista com a carga velha)."""
    from motor.carteira import checar_arquivo
    from motor.entrada import carregar_entrada, checar_base
    ent = entrada if not isinstance(entrada, (str, Path)) else carregar_entrada(entrada)
    novas, ruins, ficam = 0, [], []
    for e, destino in baixados:
        tipo = e["tipo"]
        if tipo not in TIPOS_CARTEIRA:
            ficam.append((e, destino))
            continue
        carga = tipo in ("base", "incremental")
        novas += carga
        if carga:
            ausentes = checar_base(destino, (ent.incremental or ent.base) if tipo == "incremental" else ent.base)
        else:
            ausentes = checar_arquivo(destino, getattr(ent, tipo), tipo)
        if not ausentes:
            ficam.append((e, destino))
            continue
        rej = destino.parent / "rejeitados" / destino.name
        rej.parent.mkdir(parents=True, exist_ok=True)
        destino.replace(rej)
        precisa = {"retirada": "cliente ou contrato", "acordo": "cliente ou contrato, parcela e vencimento",
                   "baixa": "cliente ou contrato, data e valor do pagamento"}.get(tipo, "")
        motivo = (f"não reconheci as colunas deste arquivo de {tipo} (precisa de {precisa}): confira e envie "
                  f"de novo" if ausentes == ["layout"] else
                  f"arquivo sem as colunas obrigatórias {ausentes}: confira o arquivo e envie de novo")
        sb.atualizar("envios", {"id": f"eq.{e['id']}"}, {"status": "erro", "relatorio": {"erro": motivo}})
        out(f"  envio {e['id']} ({e['nome_original']}): {motivo}")
        ruins.append(e["nome_original"]) if carga else None
    if ruins and len(ruins) == novas:
        raise RuntimeError(f"carga do credor rejeitada ({', '.join(ruins)}): lista do dia não gerada")
    return ficam


def _compartilhado(pasta_emp: Path, u: dict, hoje: date) -> dict:
    """O que os OUTROS credores da empresa sabem da mesma pessoa (Hot, WhatsApp, quem acionaram).
    Fica só no Mac (pasta da empresa), nunca vai para o site.

    48h juntas, com rodízio: não aciona quem outro credor acionou há menos de 48h; e cede a vez
    a outro credor que ficou esperando por essa pessoa (adiado e ainda sem acionar, até 7 dias)
    se este credor foi o último a acioná-la. Assim dois credores com a mesma régua se alternam, em vez de o
    primeiro do dia ganhar sempre. Vale a foto da primeira rodada do dia: rodar de novo (carga
    nova à tarde) não muda a lista por causa de outro credor."""
    if u["id"] is None:
        return {}
    pasta = pasta_emp / "compartilhado"
    ler = lambda a: json.loads(a.read_text(encoding="utf-8")) if a.exists() else {}  # noqa: E731
    meu = ler(pasta / f"{u['codigo']}.json").get("acionados") or {}
    hot, wa, acion = set(), set(), {}
    semana = (hoje - timedelta(days=7)).isoformat()
    for arq in sorted(pasta.glob("*.json")) if pasta.exists() else []:
        if arq.stem == u["codigo"]:
            continue
        d = ler(arq)
        hot |= {tuple(x) for x in d.get("hot", [])}
        wa |= set(d.get("whatsapp", []))
        dele = d.get("acionados") or {}
        for p, dia in dele.items():
            acion[p] = max(acion.get(p, ""), dia)
        for p, dia in (d.get("esperando") or {}).items():
            if dia >= semana and p in meu and meu[p] >= dele.get(p, ""):
                acion[p] = hoje.isoformat()      # é a vez do outro credor
    foto = u["pasta"] / "estado" / "outros_credores.json"
    if foto.exists():
        f = json.loads(foto.read_text(encoding="utf-8"))
        if f.get("data") == hoje.isoformat():
            acion = f.get("acionados") or {}
    foto.parent.mkdir(parents=True, exist_ok=True)
    foto.write_text(json.dumps({"data": hoje.isoformat(), "acionados": acion}), encoding="utf-8")
    return {"hot": hot, "whatsapp": wa, "acionados": {p: date.fromisoformat(d) for p, d in acion.items()}}


def _guardar_compartilhado(pasta_emp: Path, u: dict, r: dict, hoje: date):
    if u["id"] is None:
        return
    arq = pasta_emp / "compartilhado" / f"{u['codigo']}.json"
    arq.parent.mkdir(parents=True, exist_ok=True)
    antes = json.loads(arq.read_text(encoding="utf-8")) if arq.exists() else {}
    limite, hj = (hoje - timedelta(days=7)).isoformat(), hoje.isoformat()
    comp = r.get("compartilhar") or {}
    acion = {p: d for p, d in (antes.get("acionados") or {}).items() if limite <= d < hj}   # hoje é refeito
    acion.update({p: d.isoformat() for p, d in (comp.get("acionados") or {}).items()})
    # a espera vale até o credor acionar a pessoa (até 7 dias); a de hoje é refeita
    esper = {p: d for p, d in (antes.get("esperando") or {}).items() if limite <= d < hj}
    esper.update({p: d.isoformat() for p, d in (comp.get("esperando") or {}).items()})
    esper = {p: d for p, d in esper.items() if acion.get(p, "") < d}
    tmp = arq.with_suffix(".tmp")
    tmp.write_text(json.dumps({"hot": sorted(list(x) for x in comp.get("hot", ())),
                               "whatsapp": sorted(comp.get("whatsapp", ())), "acionados": acion,
                               "esperando": esper}), encoding="utf-8")
    tmp.replace(arq)


def _layout_automatico(sb: Supabase, pasta: Path, baixados, out):
    """Empresa sem empresas/<slug>.json: o MotorCob reconhece as colunas pelos arquivos (motor/detectar).
    Grava o que entendeu em config/entrada_automatica.json e devolve (Entrada, alerta para conferir).
    Carga que não dá para entender vai para rejeitados/ com a lista das colunas no envio."""
    from motor.detectar import entrada_automatica, frase
    from motor.ingestao import LayoutInvalido
    try:
        entrada, info = entrada_automatica(pasta)
    except LayoutInvalido as ex:
        for e, destino in baixados:
            if e["tipo"] in ("base", "incremental") and destino.exists():
                rej = destino.parent / "rejeitados" / destino.name
                rej.parent.mkdir(parents=True, exist_ok=True)
                destino.replace(rej)
                sb.atualizar("envios", {"id": f"eq.{e['id']}"},
                             {"status": "erro", "relatorio": {"erro": f"não reconheci as colunas: {ex}"}})
        raise RuntimeError(f"não reconheci as colunas da carga: {ex}") from None
    (pasta / "config").mkdir(parents=True, exist_ok=True)
    (pasta / "config" / "entrada_automatica.json").write_text(
        json.dumps(info["layout"], ensure_ascii=False, indent=1), encoding="utf-8")
    aviso = frase(info["entendido"])
    out(f"  {aviso}")
    return entrada, aviso


def _rodar_credor(sb: Supabase, emp: dict, u: dict, data: date, baixados, entrada, out, pasta_emp: Path):
    import rodar_dia
    eid, slug, cid, pasta = emp["id"], emp["slug"], u["id"], u["pasta"]
    nome = slug if cid is None else f"{slug}/{u['codigo']}"
    if cid is not None:
        out(f"  CREDOR {u['codigo']} ({u.get('nome') or ''})")
    execucao = sb.inserir("execucoes", {"empresa_id": eid, **_c(cid), "data_ref": data.isoformat(),
                                        "status": "rodando"}, retornar=True)[0]
    try:
        tem = lambda d: (pasta / d).exists() and any(p.is_file() for p in (pasta / d).iterdir())  # noqa: E731
        aviso_layout = None
        if entrada is None and any(tem(d) for d in PASTAS_CARTEIRA + ("ocorrencias", "enriquecimento")):
            entrada, aviso_layout = _layout_automatico(sb, pasta, baixados, out)
        if entrada is not None:
            baixados = _separar_cargas_ruins(sb, entrada, pasta, baixados, out)
        rel_base = None
        if tem("bruto") or tem("incremental"):
            rel_base = rodar_dia.preparar_carteira(entrada, pasta)
            out(f"  carteira: {rel_base['clientes']} clientes, {rel_base['na_carga']} em cobrança, "
                f"{rel_base['contatos']} contatos")
            # só nomes e tipos das colunas, para o site montar as regras de cluster
            sb.atualizar("empresas", {"id": f"eq.{eid}"},
                         {"colunas_base": {"colunas": rel_base["colunas"], "atualizado_em": agora()}})
        rel_enriq = None
        if tem("enriquecimento") and (pasta / "base" / "contatos.csv").exists():
            rel_enriq = rodar_dia.preparar_enriquecimento(entrada, pasta / "enriquecimento", pasta / "base")
        n_aplicadas = aplicar_sugestoes(sb, eid, out, cid)
        clusters = baixar_clusters(sb, pasta, eid, cid)
        estrategias = _baixar(sb, pasta, "estrategias", eid, "id.asc")
        if u.get("estrategia_id"):   # estratégia padrão do credor
            estrategias = [{**e, "padrao": e.get("id") == u["estrategia_id"]} for e in estrategias]
            (pasta / "config" / "estrategias.json").write_text(json.dumps(estrategias, ensure_ascii=False, indent=1),
                                                               encoding="utf-8")
        canais = _baixar(sb, pasta, "canais_empresa", eid, "canal.asc")
        for obrig in ("base/clientes.csv", "base/contatos.csv"):
            if not (pasta / obrig).exists():
                raise RuntimeError(f"falta a base de {nome}: envie a carga geral pelo site (Enviar arquivos)")
        (pasta / "retornos").mkdir(parents=True, exist_ok=True)
        parcelas = pasta / "base" / "parcelas.csv"
        acoes, portal = pasta / "acoes" / "acoes.csv", pasta / "logs" / "portal.csv"
        tem_portal = acoes.exists() and portal.exists()
        r = rodar_dia.rodar_dia(pasta / "base" / "clientes.csv", pasta / "base" / "contatos.csv", pasta / "retornos",
                                data, RAIZ / "layouts", parcelas if parcelas.exists() else None,
                                acoes if tem_portal else None, portal if tem_portal else None,
                                pasta / "estado", pasta / "saida", out=out,
                                ocorrencias=pasta / "ocorrencias" if entrada else None, entrada=entrada,
                                clusters=clusters, atributos=pasta / "base" / "atributos.csv",
                                estrategias=estrategias, canais=canais,
                                compartilhado=_compartilhado(pasta_emp, u, data))
        _guardar_compartilhado(pasta_emp, u, r, data)

        if aviso_layout:
            r["alertas"] = [aviso_layout] + r["alertas"]
        sb.inserir("estado_cliente", _linhas_estado(r["estados"], eid, r.get("clientes"), data, cid),
                   conflito="empresa_id,credor_id,id_cliente" if cid is not None else "empresa_id,id_cliente")
        n_trilha = publicar_trilha(sb, pasta / "estado", eid, cid)
        sb.apagar("fila_dia", {"empresa_id": f"eq.{eid}", **_fc(cid), "data": f"eq.{data.isoformat()}"})
        fila = _linhas_fila(r["fila"], data, eid, cid)
        if fila:
            sb.inserir("fila_dia", fila)
        prefixo = f"{slug}/{data.isoformat()}" + (f"/{u['codigo']}" if cid is not None else "")
        n_arq = publicar_arquivos(sb, r["saida"], prefixo)
        n_sug = publicar_personas(sb, eid, r, cid)
        publicar_acoes(sb, eid, r, cid)
        for e, destino in baixados:
            status, rel = _relatorio_envio(e, destino, r, rel_base, rel_enriq)
            sb.atualizar("envios", {"id": f"eq.{e['id']}"}, {"status": status, "relatorio": rel})

        estados = Counter(e.estado for e in r["estados"].values())
        resumo = {"clientes": len(r["estados"]), "estados": dict(estados),
                  "clusters": dict(Counter(e.cluster_atual for e in r["estados"].values())),
                  "fila": dict(Counter(l["canal"] for l in fila if not l["reserva"])),
                  "reserva": sum(l["reserva"] for l in fila), "enriquecimento": len(r["enriquecimento"]),
                  "dias_processados": r["dias_processados"], "trilha_enviada": n_trilha,
                  "ocorrencias": sum(x.aceitas for x in r["relatorios_ocorrencia"]),
                  "na_carga": r.get("na_carga"), "contatos": r.get("contatos_status") or {},
                  "personas": len(r.get("personas") or []), "sugestoes_novas": n_sug,
                  "sugestoes_aplicadas": n_aplicadas,
                  "acoes_ontem": _acoes_de(r.get("acoes"), data - timedelta(days=1)),
                  "quarentena": len(r["quarentena"]), "sem_layout": r["sem_layout"], "arquivos": n_arq}
        if cid is not None:
            resumo["credor"] = u["codigo"]
        if rel_base:
            resumo["carteira"] = {"em_cobranca": rel_base["na_carga"], "fora": rel_base["fora_da_carga"],
                                  "retirados": rel_base.get("retirados") or {},
                                  "pagamentos": rel_base.get("pagamentos") or {},
                                  "quitados": len(rel_base.get("quitados") or {}),
                                  "acordos": rel_base.get("acordos", 0),
                                  "acordos_abertos": rel_base.get("acordos_abertos", 0),
                                  "parcelas_pagas": rel_base.get("parcelas_pagas", 0)}
        sb.atualizar("execucoes", {"id": f"eq.{execucao['id']}"},
                     {"status": "ok", "terminada_em": agora(), "resumo": resumo, "alertas": r["alertas"]})
        out(f"  nuvem: {len(r['estados'])} clientes · {n_trilha} eventos de trilha · {len(fila)} linhas de fila "
            f"· {n_arq} arquivos publicados")
        return {"execucao": execucao["id"], "resumo": resumo, "rodar": r}
    except Exception as ex:
        sb.atualizar("execucoes", {"id": f"eq.{execucao['id']}"},
                     {"status": "erro", "terminada_em": agora(), "erro": str(ex)[:2000]})
        raise


def sincronizar_empresa_dia(sb: Supabase, dados: Path, emp: dict, data: date, out=print,
                            pasta_empresas: Path = PASTA_EMPRESAS, so_credores: set | None = None):
    """Roda o dia de cada credor ativo da empresa (ou só os de so_credores). Um credor com erro
    não impede os outros. Sem credores no banco: a empresa inteira, como antes."""
    pasta_emp = pasta_empresa(dados, emp)
    entrada = arquivo_entrada(emp, pasta_empresas)
    out(f"EMPRESA {emp['slug']}")
    unidades = credores(sb, dados, emp)
    if so_credores is not None and unidades[0]["id"] is not None:
        unidades = [u for u in unidades if u["id"] in so_credores]
    baixados = baixar_entradas(sb, dados, emp["id"], out, unidades,
                               so_credores if unidades and unidades[0]["id"] is not None else None)
    if len(unidades) == 1 and unidades[0]["id"] is None:   # banco sem credores
        return _rodar_credor(sb, emp, unidades[0], data, baixados.get(None, []), entrada, out, pasta_emp)
    res, erros = {}, []
    for u in unidades:
        try:
            res[u["codigo"]] = _rodar_credor(sb, emp, u, data, baixados.get(u["id"], []), entrada, out, pasta_emp)
        except Exception as ex:  # noqa: BLE001 — um credor com erro não para os outros
            out(f"  ERRO no credor {u['codigo']}: {ex}")
            res[u["codigo"]] = {"erro": str(ex)}
            erros.append(f"{u['codigo']}: {ex}")
    if erros and len(erros) == len(unidades):
        raise RuntimeError("; ".join(erros))
    saida = {"credores": res}
    if len(res) == 1:
        saida.update(next(iter(res.values())))
    if erros:
        saida["erro"] = "; ".join(erros)
    return saida


def sincronizar_dia(dados: Path, data: date, sb: Supabase | None = None, out=print, empresa: str | None = None,
                    pasta_empresas: Path = PASTA_EMPRESAS):
    """Roda o dia para cada empresa ativa. Retorna {slug: resultado}; erros ficam em resultado['erro']."""
    sb = sb or Supabase(*carregar_config(dados))
    resultados = {}
    for emp in empresas_ativas(sb, empresa):
        try:
            resultados[emp["slug"]] = sincronizar_empresa_dia(sb, dados, emp, data, out, pasta_empresas)
        except Exception as ex:  # noqa: BLE001 — uma empresa com erro não para as outras
            out(f"  ERRO na empresa {emp['slug']}: {ex}")
            resultados[emp["slug"]] = {"erro": str(ex)}
    if not resultados:
        out("Nenhuma empresa ativa no site: cadastre em Empresas.")
    return resultados


def _pendentes_base(sb: Supabase) -> dict:
    """{empresa_id: {credor_id: maior id de envio da carteira pendente}} (carga, retirada, acordo, baixa)."""
    por = {}
    tipos = ",".join(TIPOS_CARTEIRA)
    for e in sb.selecionar("envios", {"status": "eq.pendente", "tipo": f"in.({tipos})"}) or []:
        c = por.setdefault(e["empresa_id"], {})
        c[e.get("credor_id")] = max(c.get(e.get("credor_id"), 0), int(e["id"]))
    return por


def _vigia_arq(dados: Path, emp: dict) -> Path:
    return pasta_empresa(dados, emp) / "estado" / "vigia.json"


def empresas_com_carga_nova(dados: Path, sb: Supabase, empresa: str | None = None) -> list[dict]:
    """Empresas ativas com arquivo da carteira pendente que ainda não falhou nesse mesmo arquivo.
    Cada uma vem com "_credores": os credores que têm arquivo novo."""
    pend = _pendentes_base(sb)
    saida = []
    for emp in empresas_ativas(sb, empresa) if pend else []:
        if emp["id"] not in pend:
            continue
        arq = _vigia_arq(dados, emp)
        falhou = json.loads(arq.read_text(encoding="utf-8")).get("falhou_ate", {}) if arq.exists() else {}
        if not isinstance(falhou, dict):
            falhou = {"None": falhou}
        novos = {c for c, i in pend[emp["id"]].items() if i > falhou.get(str(c), 0)}
        if novos:
            saida.append({**emp, "_credores": novos})
    return saida


def vigiar(dados: Path, data: date, sb: Supabase | None = None, out=print, empresa: str | None = None,
           pasta_empresas: Path = PASTA_EMPRESAS):
    """Roda o dia dos credores que acabaram de receber arquivo da carteira. Retorna {slug: resultado}."""
    sb = sb or Supabase(*carregar_config(dados))
    resultados = {}
    for emp in empresas_com_carga_nova(dados, sb, empresa):
        alvo = emp.pop("_credores")
        pend = _pendentes_base(sb).get(emp["id"], {})
        out(f"ARQUIVO NOVO DA CARTEIRA: {emp['slug']} — gerando a lista de {data:%d/%m/%Y}")
        arq = _vigia_arq(dados, emp)
        falhou = json.loads(arq.read_text(encoding="utf-8")).get("falhou_ate", {}) if arq.exists() else {}
        falhou = falhou if isinstance(falhou, dict) else {}
        try:
            r = sincronizar_empresa_dia(sb, dados, emp, data, out, pasta_empresas, so_credores=alvo)
            resultados[emp["slug"]] = r
            erro = r.get("erro")
        except Exception as ex:  # noqa: BLE001 — uma empresa com erro não para as outras
            out(f"  ERRO na empresa {emp['slug']}: {ex}")
            resultados[emp["slug"]] = {"erro": str(ex)}
            erro = str(ex)
        if erro:   # não repete o mesmo erro: só tenta de novo quando chegar outro arquivo
            falhou.update({str(c): pend.get(c, 0) for c in alvo})
            arq.parent.mkdir(parents=True, exist_ok=True)
            arq.write_text(json.dumps({"falhou_ate": falhou, "erro": erro[:500]}), encoding="utf-8")
        elif arq.exists():
            arq.unlink()
    return resultados


# ------------------------------------------------------------------ comitê
def sincronizar_comite(dados: Path, mes: str, sb: Supabase | None = None, out=print, empresa: str | None = None,
                       pasta_empresas: Path = PASTA_EMPRESAS):
    import relatorio
    import rodar_dia
    from motor.certificacao import certificar_contatos
    from motor.entrada import carregar_entrada, ingerir_ocorrencias
    from motor.ingestao import (carregar_carteira, carregar_clientes, carregar_layouts, carregar_parcelas,
                                ingerir_pasta)
    from motor.regua import carregar_regua

    sb = sb or Supabase(*carregar_config(dados))
    ano, m = map(int, mes.split("-"))
    inicio, fim = date(ano, m, 1), date(ano, m, calendar.monthrange(ano, m)[1])
    regua = carregar_regua()
    resultados = {}
    for emp in empresas_ativas(sb, empresa):
        for u in credores(sb, dados, emp):
            pasta, cid = u["pasta"], u["id"]
            nome = emp["slug"] if cid is None else f"{emp['slug']}/{u['codigo']}"
            if not (pasta / "estado" / "trilha.csv").exists():
                out(f"  {nome}: sem rotina rodada ainda, sem comitê")
                continue
            clientes, _ = carregar_clientes(pasta / "base" / "clientes.csv")
            contatos, pessoa_de, _ = carregar_carteira(pasta / "base" / "contatos.csv")
            eventos = ingerir_pasta(pasta / "retornos", carregar_layouts(RAIZ / "layouts"))[0]
            entrada = arquivo_entrada(emp, pasta_empresas)
            if entrada is None and (pasta / "config" / "entrada_automatica.json").exists():
                entrada = pasta / "config" / "entrada_automatica.json"
            if entrada:
                eventos += ingerir_ocorrencias(pasta / "ocorrencias", carregar_entrada(entrada), pasta / "estado")[0]
            parc = pasta / "base" / "parcelas.csv"
            parcelas = carregar_parcelas(parc)[0] if parc.exists() else {}
            estados, _ = rodar_dia.carregar_estado(pasta / "estado")
            with open(pasta / "estado" / "trilha.csv", newline="", encoding="utf-8") as f:
                trilha = list(csv.DictReader(f))
            certs = certificar_contatos([e for e in eventos if e.data <= fim], fim, contatos, pessoa_de)
            saida = pasta / "saida" / "comite" / mes
            res = relatorio.montar(clientes, estados, trilha, eventos, parcelas, regua, inicio, fim, saida, certs,
                                   out=out)
            prefixo = f"{emp['slug']}/comite/{mes}" + (f"/{u['codigo']}" if cid is not None else "")
            n_arq = publicar_arquivos(sb, saida, prefixo)
            sb.inserir("kpis", [{"empresa_id": emp["id"], **_c(cid), "inicio": inicio.isoformat(),
                                 "fim": fim.isoformat(), "safra": k["safra"], "cluster": k["cluster"], "dados": k,
                                 "gerado_em": agora()} for k in res["kpis"]],
                       conflito=("empresa_id,credor_id," if cid is not None else "empresa_id,")
                       + "inicio,fim,safra,cluster")
            out(f"  nuvem {nome}: {len(res['kpis'])} linhas de KPI e {n_arq} arquivos do comitê publicados")
            resultados[nome] = res
    return resultados


def _travar(dados: Path, esperar: bool):
    """Uma rodada por vez (agendamento diário e vigia não se atropelam). Devolve o arquivo
    travado, ou None se outra rodada estiver em andamento e esperar=False."""
    import fcntl
    arq = Path(dados) / "logs" / ".rodando.lock"
    arq.parent.mkdir(parents=True, exist_ok=True)
    f = open(arq, "w")
    try:
        fcntl.flock(f, fcntl.LOCK_EX | (0 if esperar else fcntl.LOCK_NB))
    except BlockingIOError:
        f.close()
        return None
    return f


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("modo", choices=["dia", "comite", "vigiar"])
    ap.add_argument("--checar", action="store_true", help="vigiar: só diz se há carga nova (saída 0) ou não (3)")
    ap.add_argument("--dados", type=Path, default=Path(os.environ.get("MOTORCOB_DADOS", Path.home() / "MotorCob-dados")))
    ap.add_argument("--data", type=date.fromisoformat, default=date.today())
    ap.add_argument("--mes", help="AAAA-MM (modo comite)")
    ap.add_argument("--empresa", help="slug de uma empresa (padrão: todas as ativas)")
    a = ap.parse_args()
    try:
        if a.modo == "vigiar" and a.checar:
            sys.exit(0 if empresas_com_carga_nova(a.dados, Supabase(*carregar_config(a.dados)), a.empresa) else 3)
        trava = _travar(a.dados, esperar=a.modo != "vigiar")
        if trava is None:
            return   # outra rodada em andamento: a vigia tenta de novo daqui a pouco
        if a.modo == "vigiar":
            res = vigiar(a.dados, a.data, empresa=a.empresa)
            if any("erro" in r for r in res.values()):
                sys.exit(1)
        elif a.modo == "dia":
            res = sincronizar_dia(a.dados, a.data, empresa=a.empresa)
            if any("erro" in r for r in res.values()):
                sys.exit(1)
        else:
            if not a.mes:
                ap.error("informe --mes AAAA-MM")
            sincronizar_comite(a.dados, a.mes, empresa=a.empresa)
    except SystemExit:
        raise
    except Exception:
        traceback.print_exc()
        sys.exit(1)


if __name__ == "__main__":
    main()
