"""Campanhas da Mensageria com lista própria (sem depender da Orquestração).

1. analisar: o cliente sobe a lista em QUALQUER layout (CSV/TXT com ; , | ou tab, ou XLSX). O motor lê, identifica
   as colunas (telefone, e-mail, ID, CPF, data, valor, texto), sugere o layout e grava no banco só o NOME e o TIPO
   de cada coluna e quantas linhas vieram. O cliente confirma o layout no site (confirmar_layout).
2. preparar: na hora agendada, aplica o layout, higieniza (formato do telefone/e-mail, repetidos, lista de
   bloqueio, descadastro; favoritos passam direto pela validação de formato), preenche as variáveis do template
   e monta o lote. Contato, texto e variáveis ficam só no servidor (lote_<id>.json); no banco, o ID do cliente, o
   destino mascarado e o hash. ID que parece CPF vira um código (o banco nunca guarda CPF).
3. concluir: campanha com todos os lotes terminados vira concluída.
O envio é o mesmo dos lotes da Lista do dia (mensageria.processar).
"""
import csv
import hashlib
import io
import json
import os
import re
import secrets
import unicodedata
from datetime import datetime, timedelta
from pathlib import Path

from nuvem import mensageria as M

DDDS = {11, 12, 13, 14, 15, 16, 17, 18, 19, 21, 22, 24, 27, 28, 31, 32, 33, 34, 35, 37, 38, 41, 42, 43, 44, 45, 46,
        47, 48, 49, 51, 53, 54, 55, 61, 62, 63, 64, 65, 66, 67, 68, 69, 71, 73, 74, 75, 77, 79, 81, 82, 83, 84, 85, 86,
        87, 88, 89, 91, 92, 93, 94, 95, 96, 97, 98, 99}
EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[a-z]{2,}$", re.I)
VARIAVEL = re.compile(r"\{\{\s*([a-zA-Z0-9_]+)\s*\}\}")
MAX_LINHAS = 500_000
URL = re.compile(r"https?://[^\s\"'<>]+")
BASE62 = "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz"


def dominio_links() -> str:
    """Endereço do encurtador: o domínio curto (MOTORCOB_LINK_DOMINIO, ex.: mcb.one) ou, sem ele, o endereço direto
    da função l do Supabase (funciona igual, só é mais longo)."""
    curto = os.environ.get("MOTORCOB_LINK_DOMINIO", "").strip().strip("/")
    if curto:
        return curto.removeprefix("https://").removeprefix("http://")
    from urllib.parse import urlparse
    return f"{urlparse(os.environ.get('MOTORCOB_SUPABASE_URL', '')).netloc or 'localhost'}/functions/v1/l"


def encurtar(valor, links: dict):
    """Troca cada link (em textos e no conteúdo do RCS) por um endereço curto único. links: {codigo: url}."""
    if isinstance(valor, str):
        def troca(m):
            url = m.group(0).rstrip(".,;:!?)")
            resto = m.group(0)[len(url):]
            codigo = "".join(secrets.choice(BASE62) for _ in range(8))
            links[codigo] = url
            return f"https://{dominio_links()}/{codigo}{resto}"
        return URL.sub(troca, valor)
    if isinstance(valor, list):
        return [encurtar(v, links) for v in valor]
    if isinstance(valor, dict):
        return {k: encurtar(v, links) for k, v in valor.items()}
    return valor
RETENCAO_DIAS = 90


# ------------------------------------------------------------------ leitura e identificação do layout
def ler_tabela(conteudo: bytes, nome: str) -> tuple[list[str], list[list[str]]]:
    """(cabeçalho, linhas) de CSV/TXT (separador detectado) ou XLSX (primeira aba)."""
    if nome.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True).worksheets[0]
        linhas = [["" if v is None else (v.strftime("%Y-%m-%d") if hasattr(v, "strftime") else
                   (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v))).strip() for v in r]
                  for r in ws.iter_rows(values_only=True)]
    else:
        try:
            texto = conteudo.decode("utf-8-sig")
        except UnicodeDecodeError:
            texto = conteudo.decode("latin-1")
        amostra = texto[:20000]
        sep = max([";", ",", "\t", "|"], key=lambda s: amostra.splitlines()[0].count(s) if amostra else 0)
        linhas = [[c.strip() for c in r] for r in csv.reader(io.StringIO(texto), delimiter=sep)]
    linhas = [r for r in linhas if any(c for c in r)]
    if not linhas:
        raise ValueError("arquivo vazio")
    cab = [c or f"coluna_{i + 1}" for i, c in enumerate(linhas[0])]
    if _parece_dado(cab):           # sem cabeçalho: colunas numeradas
        cab, corpo = [f"coluna_{i + 1}" for i in range(len(cab))], linhas
    else:
        corpo = linhas[1:]
    if len(corpo) > MAX_LINHAS:
        raise ValueError(f"lista acima de {MAX_LINHAS:,} linhas: divida em mais campanhas".replace(",", "."))
    return cab, [r + [""] * (len(cab) - len(r)) for r in corpo]


def _parece_dado(cab: list[str]) -> bool:
    return sum(1 for c in cab if telefone(c) or EMAIL.match(c) or re.fullmatch(r"\d{5,}", c)) >= max(1, len(cab) // 2)


def telefone(valor: str) -> str | None:
    """Celular brasileiro normalizado (DDD + 9 dígitos) ou None."""
    d = re.sub(r"\D", "", valor or "")
    if re.fullmatch(r"55\d{11}", d):
        d = d[2:]
    if len(d) == 11 and int(d[:2]) in DDDS and d[2] == "9":
        return d
    return None


def cpf_valido(valor: str) -> bool:
    d = re.sub(r"\D", "", valor or "")
    if len(d) != 11 or d == d[0] * 11:
        return False
    for n in (9, 10):
        soma = sum(int(d[i]) * (n + 1 - i) for i in range(n))
        if (soma * 10 % 11) % 10 != int(d[n]):
            return False
    return True


def slug(nome: str) -> str:
    s = unicodedata.normalize("NFKD", nome).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "_", s).strip("_")[:40] or "coluna"


def _tipo(valores: list[str], nome: str) -> str:
    v = [x for x in valores if x]
    if not v:
        return "vazia"
    n = len(v)
    taxa = lambda f: sum(1 for x in v if f(x)) / n  # noqa: E731
    if taxa(telefone) >= 0.6:
        return "telefone"
    if taxa(lambda x: bool(EMAIL.match(x))) >= 0.6:
        return "email"
    if taxa(cpf_valido) >= 0.6 or re.search(r"\bcpf\b|cpf_cnpj|documento", nome, re.I):
        return "cpf"
    if taxa(lambda x: bool(re.fullmatch(r"\d{4}-\d{2}-\d{2}|\d{2}/\d{2}/\d{4}", x))) >= 0.6:
        return "data"
    if taxa(lambda x: bool(re.fullmatch(r"(R\$\s*)?-?[\d.]+(,\d{1,2})?|-?\d+(\.\d{1,2})?", x))) >= 0.6:
        return "valor" if re.search(r"valor|saldo|divida|d[ée]bito|parcela|pre[cç]o|r\$", nome, re.I) else "numero"
    return "texto"


def detectar_layout(cab: list[str], linhas: list[list[str]]) -> dict:
    """Tipos das colunas e o layout sugerido. Não guarda nenhum valor da lista."""
    amostra = linhas[:2000]
    colunas, usados = [], set()
    for i, nome in enumerate(cab):
        valores = [r[i] for r in amostra]
        preenchidas = sum(1 for x in valores if x)
        distintos = len(set(x for x in valores if x))
        colunas.append({"nome": nome, "tipo": _tipo(valores, nome),
                        "preenchidas_pct": round(100 * preenchidas / max(1, len(valores))),
                        "unica": preenchidas > 0 and distintos == preenchidas, "variavel": slug(nome)})
    sug, avisos = {"variaveis": {}}, []

    def melhor(tipo):
        cands = [c for c in colunas if c["tipo"] == tipo and c["nome"] not in usados]
        return max(cands, key=lambda c: c["preenchidas_pct"])["nome"] if cands else None

    for chave in ("telefone", "email"):
        sug[chave] = melhor(chave)
        if sug[chave]:
            usados.add(sug[chave])
    por_nome = [c for c in colunas if c["nome"] not in usados and c["tipo"] not in ("telefone", "email", "cpf")
                and re.search(r"^(id|codigo|c[oó]d|contrato|cliente|matr[ií]cula)", c["nome"], re.I)]
    unicos = [c for c in colunas if c["nome"] not in usados and c["unica"] and c["tipo"] in ("numero", "texto")]
    sug["id_cliente"] = (por_nome or unicos or [{"nome": None}])[0]["nome"]
    if sug["id_cliente"]:
        usados.add(sug["id_cliente"])
    for c in colunas:
        if c["nome"] not in usados and c["tipo"] not in ("vazia", "telefone", "email"):
            sug["variaveis"][c["variavel"]] = c["nome"]
    if any(c["tipo"] == "cpf" for c in colunas):
        avisos.append("A lista tem CPF. Ele pode ser usado como variável do texto, mas nunca é gravado no banco.")
    if not sug["telefone"] and not sug["email"]:
        avisos.append("Não achamos coluna de telefone celular nem de e-mail: escolha a coluna do contato.")
    if not sug["id_cliente"]:
        avisos.append("Sem coluna de ID: cada linha recebe um código (L1, L2…).")
    return {"colunas": colunas, "sugestao": sug, "linhas": len(linhas), "avisos": avisos}


# ------------------------------------------------------------------ pastas e arquivos locais
def pasta_campanhas(dados: Path, slug_empresa: str) -> Path:
    return Path(dados) / "empresas" / slug_empresa / "mensageria_campanhas"


def _empresa(sb, eid) -> dict:
    return (sb.selecionar("empresas", {"id": f"eq.{eid}"}) or [{}])[0]


# ------------------------------------------------------------------ 1. analisar
def analisar(sb, dados: Path, out=print) -> int:
    n = 0
    for c in sb.selecionar("mensageria_campanhas", {"situacao": "eq.analisando"}, ordem="id.asc") or []:
        if not c.get("arquivo_caminho"):
            continue                                   # o site ainda não terminou de subir o arquivo
        filtro = {"id": f"eq.{c['id']}"}
        try:
            conteudo = sb.baixar("entradas", c["arquivo_caminho"])
            cab, linhas = ler_tabela(conteudo, c.get("arquivo_nome") or c["arquivo_caminho"])
            lay = detectar_layout(cab, linhas)
            sb.atualizar("mensageria_campanhas", filtro, {"layout_sugerido": lay, "carregados": len(linhas),
                                                          "situacao": "aguardando_layout", "erro": None})
            out(f"campanha {c['id']}: {len(linhas)} linhas, layout sugerido")
            n += 1
        except Exception as e:  # noqa: BLE001 — arquivo ruim vira erro na campanha, não derruba o plantão
            sb.atualizar("mensageria_campanhas", filtro, {"situacao": "erro", "erro": f"não deu para ler: {e}"[:500]})
    return n


# ------------------------------------------------------------------ 2. preparar
def preencher(valor, variaveis: dict):
    """Troca {{var}} pelos valores da linha em textos e em qualquer texto dentro do conteúdo do RCS."""
    if isinstance(valor, str):
        return VARIAVEL.sub(lambda m: variaveis.get(m.group(1).lower(), ""), valor)
    if isinstance(valor, list):
        return [preencher(v, variaveis) for v in valor]
    if isinstance(valor, dict):
        return {k: preencher(v, variaveis) for k, v in valor.items()}
    return valor


def texto_rcs(conteudo: dict) -> str:
    """Texto principal do RCS (também usado no fallback por SMS): o texto, ou título e descrição do (primeiro) card."""
    if conteudo.get("texto"):
        return conteudo["texto"].strip()
    card = conteudo.get("card") or (conteudo.get("cards") or [{}])[0]
    return " — ".join(x for x in (card.get("titulo"), card.get("descricao")) if x).strip()


def _id_seguro(valor: str, linha: int) -> str:
    v = (valor or "").strip()
    if not v:
        return f"L{linha}"
    if cpf_valido(v):                                  # o banco nunca guarda CPF
        return "H" + hashlib.sha256(re.sub(r"\D", "", v).encode()).hexdigest()[:12]
    return v[:80]


def listas_contato(sb, empresa_id, canal: str) -> tuple[set, set]:
    """(bloqueio, favoritos) em hashes, do canal e de 'todos'."""
    linhas = M._tabela_nova(lambda: sb.selecionar("mensageria_listas_contato", {"empresa_id": f"eq.{empresa_id}"})
                            or [], [])
    sel = lambda t: {x["destino_hash"] for x in linhas if x["tipo"] == t and x["canal"] in ("todos", canal)}  # noqa
    return sel("bloqueio"), sel("favorito")


def preparar(sb, dados: Path, agora: datetime | None = None, out=print) -> list[dict]:
    agora = (agora or datetime.now(M.SP)).astimezone(M.SP)
    criados = []
    for c in sb.selecionar("mensageria_campanhas", {"situacao": "eq.agendada"}, ordem="id.asc") or []:
        quando = c.get("agendada_para")
        if quando and datetime.fromisoformat(str(quando).replace("Z", "+00:00")) > agora:
            continue
        try:
            criados.append(_montar(sb, dados, c, agora, out))
        except Exception as e:  # noqa: BLE001
            sb.atualizar("mensageria_campanhas", {"id": f"eq.{c['id']}"},
                         {"situacao": "erro", "erro": f"não deu para montar o envio: {e}"[:500]})
    return criados


def _montar(sb, dados: Path, c: dict, agora: datetime, out) -> dict:
    eid, canal, lay = c["empresa_id"], c["canal"], c.get("layout") or {}
    emp = _empresa(sb, eid)
    tpl = (sb.selecionar("mensageria_templates", {"id": f"eq.{c['template_id']}"}) or [None])[0] \
        if c.get("template_id") else None
    if canal in ("sms", "rcs") and not tpl:
        raise ValueError("campanha sem template")
    modelo_email = None
    if canal == "email":
        modelo_email = (sb.selecionar("mensageria_modelos_email", {"empresa_id": f"eq.{eid}"}) or [None])[0]
        if not modelo_email:
            raise ValueError("cadastre o modelo de e-mail")
    fb_tpl = None
    if canal == "rcs" and c.get("fallback_template_id"):
        fb_tpl = (sb.selecionar("mensageria_templates", {"id": f"eq.{c['fallback_template_id']}"}) or [None])[0]
    cab, linhas = ler_tabela(sb.baixar("entradas", c["arquivo_caminho"]), c.get("arquivo_nome") or c["arquivo_caminho"])
    idx = {n: i for i, n in enumerate(cab)}
    col = lambda chave: idx.get(lay.get(chave)) if lay.get(chave) else None  # noqa: E731
    i_dest, i_id = col("email" if canal == "email" else "telefone"), col("id_cliente")
    if i_dest is None:
        raise ValueError("coluna do contato não encontrada no arquivo")
    var_cols = {v: idx[n] for v, n in (lay.get("variaveis") or {}).items() if n in idx}
    from nuvem.atendimento import banco_radar
    rastrear = bool(c.get("rastrear_links") or (tpl or {}).get("rastrear_links")) and banco_radar(sb)
    links_por_linha = {}
    bloqueio, favoritos = listas_contato(sb, eid, canal)
    descad = {d["destino_hash"] for d in sb.selecionar("mensageria_descadastros", {
        "empresa_id": f"eq.{eid}", "canal": f"eq.{canal}"}) or []}
    cfg = (sb.selecionar("mensageria_config", {"empresa_id": f"eq.{eid}"}) or [{}])[0]
    modo = cfg.get("modo") or "simulado"
    canais_cfg = {x.get("canal"): x for x in sb.selecionar("canais_empresa", {"empresa_id": f"eq.{eid}"}) or []}
    linhas_db, locais, vistos = [], {}, set()
    cont = {"invalidos": 0, "duplicados": 0, "bloqueados": 0}
    for n, r in enumerate(linhas, start=1):
        bruto = r[i_dest].strip()
        idc = _id_seguro(r[i_id] if i_id is not None else "", n)
        h = M.hash_destino(canal, bruto)
        fav = h in favoritos
        dest = (bruto.lower() if EMAIL.match(bruto) else None) if canal == "email" else (telefone(bruto) or
                                                                                         (re.sub(r"\D", "", bruto) if fav else None))
        if not dest:
            cont["invalidos"] += 1
            linhas_db.append(_msg(idc, canal, bruto or "—", h, "bloqueada", "contato inválido na higienização"))
            continue
        if h in vistos:
            cont["duplicados"] += 1
            continue
        vistos.add(h)
        variaveis = {v: r[i].strip() for v, i in var_cols.items()}
        motivo, texto, conteudo, item = None, None, None, canal
        if h in bloqueio:
            motivo = "lista de bloqueio"
        elif h in descad:
            motivo = "descadastrado"
        elif canal == "sms":
            texto = preencher(tpl["conteudo"].get("texto", ""), variaveis).strip()
            if rastrear:
                links_por_linha[(idc, h)] = {}
                texto = encurtar(texto, links_por_linha[(idc, h)])
            motivo = "frase acima de 160 caracteres" if len(texto) > M.LIMITE_SMS else None if texto else "texto vazio"
        elif canal == "rcs":
            conteudo = preencher(tpl["conteudo"], variaveis)
            if rastrear:
                links_por_linha[(idc, h)] = {}
                conteudo = encurtar(conteudo, links_por_linha[(idc, h)])
            texto = texto_rcs(conteudo)
            item = "rcs_basic" if tpl["tipo_midia"] == "simples" and len(texto) <= M.LIMITE_RCS_BASIC else "rcs_single"
        elif canal == "email":
            texto = preencher(modelo_email["corpo"], variaveis)
            if rastrear:
                links_por_linha[(idc, h)] = {}
                texto = encurtar(texto, links_por_linha[(idc, h)])
        if motivo:
            cont["bloqueados"] += 1
        linhas_db.append({**_msg(idc, canal, dest, h, "bloqueada" if motivo else "pendente", motivo), "item": item})
        if not motivo:
            loc = {"contato": dest, "texto": texto, "item": item, "favorito": fav}
            if conteudo is not None:
                loc["conteudo"] = conteudo
            if canal == "email":
                loc["assunto"] = preencher(modelo_email["assunto"], variaveis)
                loc["remetente"] = modelo_email.get("remetente_nome")
            if c.get("fallback") not in (None, "empresa"):
                loc["fallback"] = c["fallback"]
            if fb_tpl:
                loc["fallback_texto"] = preencher(fb_tpl["conteudo"].get("texto", ""), variaveis).strip()
            locais[(idc, h)] = loc
    jan = M.janela(canais_cfg.get(canal), agora.date()) or M.JANELA_PADRAO
    situacao = "agendado" if locais else "concluido"
    lote = sb.inserir("mensageria_lotes", [{
        "empresa_id": eid, "credor_id": 0, "data": agora.date().isoformat(), "canal": canal, "modo": modo,
        "situacao": situacao, "total": len(linhas_db), "bloqueadas": sum(1 for x in linhas_db if x["situacao"] == "bloqueada"),
        "invalidos": cont["invalidos"], "janela_inicio": jan[0], "janela_fim": jan[1], "campanha_id": c["id"],
        "centro_custo_id": c.get("centro_custo_id") or cfg.get("centro_custo_id"),
        "motivo": None if locais else "nada a enviar: todos bloqueados na higienização"}], retornar=True)[0]
    gravadas = sb.inserir("mensageria_mensagens", [{**x, "lote_id": lote["id"], "empresa_id": eid} for x in linhas_db],
                          retornar=True) if linhas_db else []
    links = [{"codigo": cod, "empresa_id": eid, "lote_id": lote["id"], "mensagem_id": m["id"], "campanha_id": c["id"],
              "canal": canal, "id_cliente": m["id_cliente"], "url_destino": url[:2000]}
             for m in gravadas if m["situacao"] == "pendente"
             for cod, url in links_por_linha.get((m["id_cliente"], m["destino_hash"]), {}).items()]
    if links:
        sb.inserir("mensageria_links", links)
    M._gravar_json(pasta_campanhas(dados, emp["slug"]) / f"lote_{lote['id']}.json",
                   {str(m["id"]): locais[(m["id_cliente"], m["destino_hash"])] for m in gravadas
                    if (m["id_cliente"], m["destino_hash"]) in locais})
    sb.atualizar("mensageria_campanhas", {"id": f"eq.{c['id']}"},
                 {**cont, "situacao": "em_andamento" if locais else "concluida"})
    out(f"campanha {c['id']} ({canal}): {len(locais)} para enviar, {cont['invalidos']} inválidos, "
        f"{cont['duplicados']} repetidos, {cont['bloqueados']} bloqueados")
    return lote


def _msg(idc, canal, destino, h, situacao, motivo) -> dict:
    return {"id_cliente": idc, "destino_mascara": M.mascarar(canal, destino) if destino != "—" else "—",
            "destino_hash": h, "situacao": situacao, "motivo": motivo}


# ------------------------------------------------------------------ 3. concluir e limpar
def concluir(sb) -> int:
    n = 0
    for c in sb.selecionar("mensageria_campanhas", {"situacao": "eq.em_andamento"}) or []:
        lotes = sb.selecionar("mensageria_lotes", {"campanha_id": f"eq.{c['id']}"}) or []
        if lotes and all(l["situacao"] in ("concluido", "cancelado") for l in lotes):
            sb.atualizar("mensageria_campanhas", {"id": f"eq.{c['id']}"}, {"situacao": "concluida"})
            n += 1
    return n


def limpar(sb, dados: Path, agora: datetime | None = None, out=print) -> int:
    """Apaga arquivos de lista e conteúdo local com mais de 90 dias (LGPD)."""
    agora = (agora or datetime.now(M.SP)).astimezone(M.SP)
    limite = agora - timedelta(days=RETENCAO_DIAS)
    n = 0
    for c in sb.selecionar("mensageria_campanhas", {"criado_em": f"lt.{limite.isoformat()}",
                                                    "arquivo_caminho": "not.is.null"}) or []:
        try:
            sb.apagar_arquivos("entradas", [c["arquivo_caminho"]])
        except Exception:  # noqa: BLE001 — já apagado
            pass
        sb.atualizar("mensageria_campanhas", {"id": f"eq.{c['id']}"}, {"arquivo_caminho": None})
        n += 1
    for arq in list(Path(dados).glob("empresas/*/mensageria_campanhas/lote_*.json")) + \
            list(Path(dados).glob("empresas/**/mensageria/lote_*.json")):
        if datetime.fromtimestamp(arq.stat().st_mtime, M.SP) < limite:
            arq.unlink()
            n += 1
    if n:
        out(f"mensageria: {n} arquivos com mais de {RETENCAO_DIAS} dias apagados")
    return n
