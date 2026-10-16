"""Relatórios (SQL 39): foto da carteira, exportação grande em páginas e envio agendado pela função."""
import csv
import io
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path

from nuvem import relatorios as R
from nuvem import sincronizar as S


class SbFalso:
    def __init__(self, linhas=0, pedido=None, falhar_pagina=False):
        self.linhas, self.falhar_pagina = linhas, falhar_pagina
        self.pedido = pedido
        self.rpcs, self.atualizados, self.enviados, self.avisos, self.funcoes, self.apagados = [], [], [], [], [], []

    def selecionar(self, tabela, filtros=None, ordem=None, limite=None):
        if tabela == "relatorio_exportacoes":
            return [self.pedido] if self.pedido else []
        if tabela == "relatorio_campos":
            return [{"campo": "data", "rotulo": "Data"}, {"campo": "canal", "rotulo": "Canal"}]
        return []

    def rpc(self, nome, args=None):
        self.rpcs.append((nome, args))
        if nome == "rodar_relatorio_pagina":
            if self.falhar_pagina:
                raise RuntimeError("tempo esgotado")
            ini, qtd = args["inicio"], args["qtd"]
            fim = min(self.linhas, ini + qtd)
            return [{"data": "2026-10-0%d" % (1 + i % 9), "canal": "sms", "enviados": i, "taxa_entrega": 93.5}
                    for i in range(ini, fim)]
        if nome == "tirar_foto_carteira":
            return 42
        if nome == "exportacoes_vencidas":
            return [{"id": 1, "arquivo": "7/1.xlsx"}, {"id": 2, "arquivo": None}]
        if nome == "envios_relatorio_devidos":
            return [5, {"envios_relatorio_devidos": 6}]
        return None

    def atualizar(self, tabela, filtros, valores):
        self.atualizados.append(valores)

    def enviar(self, bucket, caminho, conteudo, tipo="text/csv"):
        self.enviados.append((bucket, caminho, conteudo, tipo))

    def link_assinado(self, bucket, caminho, segundos):
        return f"https://sb.invalid/storage/v1/object/sign/{bucket}/{caminho}?token=x"

    def inserir(self, tabela, linha, **k):
        self.avisos.append(linha)

    def apagar_arquivos(self, bucket, caminhos):
        self.apagados.append((bucket, caminhos))

    def funcao(self, nome, corpo):
        self.funcoes.append((nome, corpo))
        if corpo["envio"] == 6:
            raise RuntimeError("Locaweb fora do ar")
        return {"enviado_para": ["a@x.com", "b@x.com"]}


def pedido(**k):
    p = {"id": 9, "empresa_id": 7, "pedido_por": "u-1", "nome": "Envios de outubro", "base": "funil", "colunas": None,
         "filtros": {"empresa": 7}, "agrupar": ["data", "canal"],
         "metricas": [{"funcao": "somar", "campo": "enviados", "nome": "enviados", "rotulo": "Enviadas"},
                      {"funcao": "taxa", "num": "entregues", "den": "enviados", "nome": "taxa_entrega",
                       "rotulo": "Taxa de entrega"}],
         "ordem": "data asc", "formato": "csv", "situacao": "pendente", "linhas_total": 120_000,
         "criado_em": "2026-10-15T10:00:00+00:00"}
    p.update(k)
    return p


class TestExportacao(unittest.TestCase):
    def test_csv_em_paginas_com_rotulos_e_aviso(self):
        sb = SbFalso(linhas=120_000, pedido=pedido())
        self.assertEqual(R.exportar(sb, out=lambda *a: None), 9)
        paginas = [a for n, a in sb.rpcs if n == "rodar_relatorio_pagina"]
        self.assertEqual([(a["inicio"], a["qtd"]) for a in paginas], [(0, 50_000), (50_000, 50_000), (100_000, 20_000)])
        bucket, caminho, conteudo, tipo = sb.enviados[0]
        self.assertEqual((bucket, caminho, tipo), ("relatorios", "7/9.csv", "text/csv"))
        texto = conteudo.decode("utf-8")
        self.assertTrue(texto.startswith("﻿Data;Canal;Enviadas;Taxa de entrega"))
        linhas = list(csv.reader(io.StringIO(texto.lstrip("﻿")), delimiter=";"))
        self.assertEqual(len(linhas), 120_001)
        self.assertEqual(linhas[1], ["01/10/2026", "sms", "0", "93,5"])
        final = sb.atualizados[-1]
        self.assertEqual((final["situacao"], final["linhas_feitas"]), ("pronta", 120_000))
        self.assertTrue(final["link"].startswith("https://"))
        self.assertEqual(sb.avisos[-1]["usuario_id"], "u-1")
        self.assertEqual(sb.avisos[-1]["link"], "/relatorios?exportacao=9")
        self.assertIn("120.000 linhas", sb.avisos[-1]["texto"])

    def test_excel_abre_com_datas(self):
        sb = SbFalso(linhas=3, pedido=pedido(formato="xlsx", linhas_total=3))
        R.exportar(sb, out=lambda *a: None)
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(sb.enviados[0][2])).active
        self.assertEqual([c.value for c in ws[1]], ["Data", "Canal", "Enviadas", "Taxa de entrega"])
        self.assertEqual(ws.max_row, 4)
        self.assertEqual(ws["A2"].value.date().isoformat(), "2026-10-01")

    def test_nunca_passa_de_um_milhao(self):
        sb = SbFalso(linhas=1_200_000, pedido=pedido(linhas_total=5_000_000))
        R.exportar(sb, out=lambda *a: None)
        self.assertEqual(sb.atualizados[-1]["linhas_feitas"], 1_000_000)

    def test_erro_avisa_quem_pediu(self):
        sb = SbFalso(linhas=10, pedido=pedido(), falhar_pagina=True)
        R.exportar(sb, out=lambda *a: None)
        self.assertEqual(sb.atualizados[-1]["situacao"], "erro")
        self.assertEqual(sb.avisos[-1]["tipo"], "erro")
        self.assertEqual(sb.enviados, [])

    def test_processando_recente_nao_e_pego_de_novo(self):
        agora = datetime(2026, 10, 15, 10, 10, tzinfo=timezone.utc)
        sb = SbFalso(linhas=1, pedido=pedido(situacao="processando", iniciado_em="2026-10-15T10:00:00+00:00"))
        self.assertIsNone(R.exportar(sb, out=lambda *a: None, agora=agora))
        sb = SbFalso(linhas=1, pedido=pedido(situacao="processando", iniciado_em="2026-10-15T09:00:00+00:00",
                                              linhas_total=1))
        self.assertEqual(R.exportar(sb, out=lambda *a: None, agora=agora), 9)     # travado: começa de novo


class TestFotoEnvioExpirar(unittest.TestCase):
    def test_expirar_apaga_os_arquivos(self):
        sb = SbFalso()
        self.assertEqual(R.expirar(sb, out=lambda *a: None), 2)
        self.assertEqual(sb.apagados, [("relatorios", ["7/1.xlsx"])])

    def test_envio_chama_a_funcao_e_marca_o_erro(self):
        sb = SbFalso()
        saida = []
        self.assertEqual(R.enviar(sb, out=saida.append), 1)
        self.assertEqual(sb.funcoes, [("relatorios", {"acao": "enviar_resumo", "envio": 5}),
                                      ("relatorios", {"acao": "enviar_resumo", "envio": 6})])
        self.assertIn(("marcar_envio_relatorio", {"envio": 6, "erro": "Locaweb fora do ar"}), sb.rpcs)
        self.assertTrue(any("envio 6 não saiu" in s for s in saida))

    def test_foto_uma_vez_por_dia_depois_das_20h(self):
        with tempfile.TemporaryDirectory() as d:
            sb = SbFalso()
            S._foto_carteira(sb, Path(d), datetime(2026, 10, 15, 21, 0, tzinfo=timezone.utc), out=lambda *a: None)  # 18h SP
            self.assertEqual(sb.rpcs, [])
            S._foto_carteira(sb, Path(d), datetime(2026, 10, 15, 23, 30, tzinfo=timezone.utc), out=lambda *a: None)  # 20h30
            S._foto_carteira(sb, Path(d), datetime(2026, 10, 16, 0, 30, tzinfo=timezone.utc), out=lambda *a: None)  # 21h30
            self.assertEqual([n for n, _ in sb.rpcs], ["tirar_foto_carteira"])

    def test_banco_sem_o_sql_39(self):
        class Velho:
            def rpc(self, *a, **k):
                raise RuntimeError("404")

            def selecionar(self, *a, **k):
                raise RuntimeError("404")
        self.assertIsNone(R.exportar(Velho(), out=lambda *a: None))
        self.assertEqual(R.enviar(Velho(), out=lambda *a: None), 0)
        self.assertEqual(R.expirar(Velho(), out=lambda *a: None), 0)
        self.assertEqual(R.foto(Velho(), out=lambda *a: None), 0)


if __name__ == "__main__":
    unittest.main()
