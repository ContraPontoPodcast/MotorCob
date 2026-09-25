"""Testes dos KPIs, da atribuição por régua e do relatório do comitê."""
import tempfile
import unittest
from datetime import date
from pathlib import Path

from motor.acordos import Parcela
from motor.certificacao import Evento
from motor.kpis import (LinhaDoTempo, atribuir, kpis_por_safra_cluster, matriz_migracao, realizado_por_acao,
                        sugerir_ordem_rotacao)
from motor.regua import carregar_regua

R = carregar_regua()
D = date


def t(dia, idc, antes, tag, motivo="x"):
    return {"data": dia.isoformat(), "id_cliente": idc, "tag_anterior": antes, "tag": tag, "motivo": motivo,
            "quem_marcou": "teste"}


TRILHA = [
    t(D(2026, 9, 1), "C1", "", "S260901-M1-LOC-ND-L0"),
    t(D(2026, 9, 4), "C1", "S260901-M1-LOC-ND-L0", "S260901-M1-CPA-RC-T1", "contato no RCS → CPC A"),
    t(D(2026, 9, 4), "C1", "S260901-M1-CPA-RC-T1", "S260901-M1-COL-RC", "acordo ativo (colchão)"),
]
EVENTOS = [
    Evento("C1", "11999990000", "telefone", "whatsapp", "entregue", D(2026, 9, 2), 0.30),
    Evento("C1", "11999990000", "telefone", "rcs", "interacao", D(2026, 9, 4), 0.12),
    Evento("C1", "11999990000", "telefone", "rcs", "entregue", D(2026, 9, 10), 0.12),  # já em acordo
]
PARCELAS = {"C1": [Parcela("C1", "A1", 1, D(2026, 9, 12), 500.0, D(2026, 9, 12)),
                   Parcela("C1", "A1", 2, D(2026, 10, 12), 500.0, None)]}


class TestAtribuicao(unittest.TestCase):
    def test_tentativa_vai_para_a_regua_do_estado_no_comeco_do_dia(self):
        at = atribuir(EVENTOS, LinhaDoTempo(TRILHA), R)
        self.assertEqual([(e.canal, r, c) for e, r, c in at],
                         [("whatsapp", "localizacao", False), ("rcs", "localizacao", True),
                          ("rcs", "preventivo", False)])

    def test_realizado_por_acao_liga_acordo_ao_canal_e_regua_de_origem(self):
        real = realizado_por_acao({}, TRILHA, EVENTOS, PARCELAS, R, D(2026, 9, 1), D(2026, 9, 30))
        por = {(r["regua"], r["canal"]): r for r in real}
        a = por[("localizacao", "rcs")]
        self.assertEqual((a["volume"], a["contatos"], a["acordos"], a["valor_acordado"], a["recebido"]),
                         (1, 1, 1, 1000.0, 500.0))
        self.assertEqual(por[("preventivo", "rcs")]["acordos"], 0)

    def test_matriz_de_migracao(self):
        m = matriz_migracao(TRILHA, D(2026, 9, 1), D(2026, 9, 30))
        self.assertEqual((m["ENTRADA"]["LOC"], m["LOC"]["CPA"], m["CPA"]["COL"]), (1, 1, 1))

    def test_sugestao_de_rotacao_por_custo_por_contato(self):
        real = [{"regua": "cpc", "canal": "discador", "volume": 100, "contatos": 5, "custo": 35.0},
                {"regua": "cpc", "canal": "sms", "volume": 100, "contatos": 2, "custo": 7.0}]
        ordem = sugerir_ordem_rotacao(real, R)
        self.assertEqual([o["canal"] for o in ordem], ["sms", "discador"])   # R$ 3,50 < R$ 7,00 por contato
        self.assertEqual(ordem[0]["posicao_atual"], 5)


class TestKpisSimulados(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from exemplos import simular_operacao
        cls.tmp = tempfile.TemporaryDirectory()
        cls.m = simular_operacao.simular(seed=5, dias=40, n_clientes=200, pasta=Path(cls.tmp.name), verbose=False)
        d = cls.d = cls.m["dados"]
        cls.kpis = kpis_por_safra_cluster(d["clientes"], d["estados"], d["trilha"], d["eventos"], d["parcelas"],
                                          d["regua"], d["inicio"], d["fim"], d["certs"])
        cls.real = realizado_por_acao(d["clientes"], d["trilha"], d["eventos"], d["parcelas"], d["regua"],
                                      d["inicio"], d["fim"])

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_total_bate_com_a_soma_dos_grupos(self):
        total, grupos = self.kpis[-1], self.kpis[:-1]
        for k in ("clientes", "localizados", "acordos", "quebras", "ncp_para_cpa"):
            self.assertEqual(total[k], sum(g[k] for g in grupos), k)

    def test_todo_acordo_e_todo_custo_sao_atribuidos(self):
        self.assertEqual(sum(r["acordos"] for r in self.real), self.m["acordos"])
        custo = sum(e.custo for e in self.d["eventos"])
        self.assertAlmostEqual(sum(r["custo"] for r in self.real), custo, places=1)

    def test_localizacao_por_canal_soma_o_total(self):
        total = self.kpis[-1]
        soma = sum(total[f"pct_loc_{c}"] or 0 for c in R["canais"].values())
        self.assertAlmostEqual(soma, total["pct_localizados"], places=3)

    def test_relatorio_gera_csvs_e_excel(self):
        import relatorio
        d = self.d
        with tempfile.TemporaryDirectory() as tmp:
            r = relatorio.montar(d["clientes"], d["estados"], d["trilha"], d["eventos"], d["parcelas"], d["regua"],
                                 d["inicio"], d["fim"], tmp, d["certs"], out=lambda *a: None)
            nomes = {p.name for p in Path(tmp).iterdir()}
            self.assertTrue({"kpis_safra_cluster.csv", "realizado_por_acao.csv", "benchmarks.csv"} <= nomes)
            if r["xlsx"] is None:
                self.skipTest("openpyxl não instalado")
            from openpyxl import load_workbook
            wb = load_workbook(r["xlsx"])
            self.assertEqual(wb.sheetnames, ["Leia-me", "Resumo", "Premissas", "Real x Previsto",
                                             "KPIs safra-cluster", "Migração de estados", "Benchmarks"])
            ws = wb["Real x Previsto"]
            self.assertEqual(ws["G4"].value, "=E4*F4")                   # orçado é fórmula
            self.assertTrue(str(ws["R4"].value).startswith("=IF(Q4>="))  # decisão segue o critério de ROI


if __name__ == "__main__":
    unittest.main()
